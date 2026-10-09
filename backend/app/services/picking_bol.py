from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func, select
from sqlalchemy.orm import Session, joinedload

from app.models import AuditLog, BOL, BOLItem, OutboundInventoryAllocation, PickingList, PickingListItem
from app.models.bol import BOLStatus
from app.models.picking import PickingStatus
from app.schemas.picking_bol import BOLRead, PickComplete, PickingRead
from app.services.outbound import get_ob
from app.services.cargo_bol import document_identity
from app.utils.business_time import get_business_today, to_business_date

PICK_NAMES = {0: "New", 1: "Printed", 2: "In Progress", 3: "Completed", 4: "Canceled", 5: "Exception"}
BOL_NAMES = {0: "Draft", 1: "Generated", 2: "Printed", 3: "Completed", 4: "Canceled"}
NAVY, BLUE, PALE_BLUE, PALE_YELLOW, GRID = "17324D", "1F4E78", "DDEBF7", "FFF2CC", "A6A6A6"


def number(db, model, prefix):
    root = f"{prefix}{get_business_today():%y%m%d}"
    column = model.picking_no if model is PickingList else model.bol_no
    last = db.scalar(select(func.max(column)).where(column.like(root + "%")))
    return f"{root}{(int(last[-4:]) + 1 if last else 1):04d}"


def picking_read(p):
    return PickingRead.model_validate({**p.__dict__, "status_name": PICK_NAMES[p.status], "planned_pallet_qty": sum((i.planned_pallet_qty for i in p.items), Decimal(0)), "planned_carton_qty": sum((i.planned_carton_qty for i in p.items), Decimal(0)), "picked_pallet_qty": sum((i.picked_pallet_qty for i in p.items), Decimal(0)), "picked_carton_qty": sum((i.picked_carton_qty for i in p.items), Decimal(0))})


def generate_picking(db: Session, ob_id, user_id, *, commit: bool = True):
    outbound = get_ob(db, ob_id, True)
    if outbound.status in (4, 5, 6):
        raise HTTPException(409, "Closed outbound cannot generate picking evidence")
    allocs = list(db.scalars(select(OutboundInventoryAllocation).options(joinedload(OutboundInventoryAllocation.inventory_lot)).where(OutboundInventoryAllocation.outbound_order_id == ob_id)).all())
    if not allocs:
        raise HTTPException(409, "Outbound requires allocations")
    planned = []
    for seq, allocation in enumerate(sorted(allocs, key=lambda x: (x.inventory_lot.location_id or 0, x.inventory_lot.inbound_date or date.max)), 1):
        existing = db.scalars(select(PickingListItem).join(PickingList, PickingList.id == PickingListItem.picking_list_id).where(PickingListItem.outbound_allocation_id == allocation.id, PickingList.status != PickingStatus.CANCELED)).all()
        used_pallets = sum((i.planned_pallet_qty for i in existing), Decimal(0))
        used_cartons = sum((i.planned_carton_qty for i in existing), Decimal(0))
        pallets = allocation.allocated_pallet_qty - allocation.completed_pallet_qty - used_pallets
        cartons = allocation.allocated_carton_qty - allocation.completed_carton_qty - used_cartons
        if pallets <= 0 and cartons <= 0:
            continue
        lot = allocation.inventory_lot
        planned.append(PickingListItem(**document_identity(allocation), outbound_allocation_id=allocation.id, inventory_lot_id=lot.id, location_id=lot.location_id, lot_no=lot.lot_no, container_number=lot.container_number, fc_code=lot.fc_code, marking=lot.marking, planned_pallet_qty=pallets, planned_carton_qty=cartons, planned_weight_lbs=allocation.allocated_weight_lbs - allocation.completed_weight_lbs, planned_cbm=allocation.allocated_cbm - allocation.completed_cbm, sequence_no=seq))
    if not planned:
        existing = db.scalar(select(PickingList).where(PickingList.outbound_order_id == ob_id, PickingList.status != PickingStatus.CANCELED).order_by(PickingList.id.desc()))
        if existing:
            return existing
        raise HTTPException(409, "Outbound has no quantity remaining to pick")
    p = PickingList(picking_no=number(db, PickingList, "PK"), outbound_order_id=ob_id, status=PickingStatus.NEW, created_by=user_id)
    p.items.extend(planned)
    db.add(p)
    db.flush()
    db.add(AuditLog(user_id=user_id, action="CREATE_PICKING_LIST", entity_type="PICKING", entity_id=p.id))
    if commit:
        db.commit()
    return p


def complete_picking(db: Session, pid: int, user_id: int, payload: PickComplete | None = None, *, commit: bool = True):
    p = db.get(PickingList, pid)
    if not p:
        raise HTTPException(404, "Picking list not found")
    outbound = get_ob(db, p.outbound_order_id, True)
    if outbound.status in (4, 5, 6):
        raise HTTPException(409, "Closed outbound cannot change picking evidence")
    p = db.scalar(select(PickingList).where(PickingList.id == pid).with_for_update().execution_options(populate_existing=True))
    if p.status in (3, 4):
        raise HTTPException(409, "Picking list is closed")
    for item in p.items:
        pallets = payload.picked_pallet_qty if payload and payload.picked_pallet_qty is not None else item.planned_pallet_qty
        cartons = payload.picked_carton_qty if payload and payload.picked_carton_qty is not None else item.planned_carton_qty
        if pallets < 0 or cartons < 0 or pallets > item.planned_pallet_qty or cartons > item.planned_carton_qty:
            raise HTTPException(409, "Picked quantity exceeds the picking allocation")
    for item in p.items:
        item.picked_pallet_qty = payload.picked_pallet_qty if payload and payload.picked_pallet_qty is not None else item.planned_pallet_qty
        item.picked_carton_qty = payload.picked_carton_qty if payload and payload.picked_carton_qty is not None else item.planned_carton_qty
    p.status = 3
    p.completed_at = datetime.now(UTC)
    p.completed_by = user_id
    db.add(AuditLog(user_id=user_id, action="COMPLETE_PICKING", entity_type="PICKING", entity_id=p.id))
    db.flush()
    if commit:
        db.commit()
    return p


def _text(value, fallback=""):
    return str(value).strip() if value is not None and str(value).strip() else fallback


def _transport_mode(outbound):
    raw = " ".join(filter(None, (getattr(outbound, "delivery_type", None), getattr(outbound, "truck_type", None)))).upper()
    if any(word in raw for word in ("FTL", "FULL", "TRUCKLOAD")):
        return "FTL"
    if "LTL" in raw or "LESS" in raw:
        return "LTL"
    if any(word in raw for word in ("SPD", "PARCEL", "SMALL")):
        return "SPD"
    return "LTL/FTL" if getattr(outbound, "fba_shipment", None) else "FREIGHT"


def amazon_bol_context(b):
    """Map WMS data to Amazon freight reference fields without inventing identifiers."""
    outbound = b.outbound
    fba = getattr(b, "fba_shipment", None) or getattr(outbound, "fba_shipment", None)
    load = getattr(outbound, "load", None)
    carrier = getattr(b, "carrier", None) or getattr(outbound, "carrier", None) or getattr(load, "carrier", None)
    mode = _transport_mode(outbound)
    is_fba = bool(fba or getattr(outbound, "ob_type", "") == "FBA")
    context = {
        "is_fba": is_fba,
        "mode": mode,
        "outbound_no": _text(getattr(outbound, "ob_no", None)),
        "amazon_reference_id": _text(getattr(fba, "reference_no", None)),
        "shipment_id": _text(getattr(fba, "shipment_id", None)),
        "st_number": _text(getattr(fba, "st_number", None)),
        "po_number": _text(getattr(fba, "po_number", None) or getattr(outbound, "po_number", None)),
        "appointment_reference": _text(getattr(load, "appointment_reference", None) or getattr(outbound, "appointment_reference", None)),
        # This is the imported carrier document reference, not the WMS-generated BOL number.
        "carrier_tracking": _text(getattr(outbound, "bol_reference", None)),
        "carrier_name": _text(getattr(carrier, "carrier_name", None), "To be assigned"),
        "scac": _text(getattr(carrier, "scac", None)),
        "tractor_no": _text(getattr(load, "tractor_no", None) or getattr(outbound, "truck_number", None)),
        "trailer_no": _text(getattr(load, "trailer_no", None) or getattr(outbound, "trailer_number", None)),
        "pickup_date": _text(getattr(b, "pickup_date", None)),
        "appointment_time": _text(getattr(load, "appointment_time", None) or getattr(outbound, "delivery_appointment_time", None) or getattr(b, "appointment_time", None)),
    }
    missing = []
    if is_fba and mode != "SPD":
        for label, value in (("Amazon Reference ID", context["amazon_reference_id"]), ("FBA Shipment ID", context["shipment_id"]), ("Carrier PRO / tracking / BOL", context["carrier_tracking"])):
            if not value:
                missing.append(label)
    if is_fba and (not b.ship_to_address or not b.amazon_fc_code):
        missing.append("Amazon FC destination")
    context["missing"] = missing
    context["ready"] = not missing
    return context


def bol_read(b):
    context = amazon_bol_context(b)
    return BOLRead.model_validate({**b.__dict__, "status_name": BOL_NAMES[b.status], "total_pallet_qty": sum((i.pallet_qty for i in b.items), Decimal(0)), "total_carton_qty": sum((i.carton_qty for i in b.items), Decimal(0)), "total_weight_lbs": sum((i.weight_lbs for i in b.items), Decimal(0)), "total_cbm": sum((i.cbm for i in b.items), Decimal(0)), "fc_address_missing": bool(b.amazon_fc_code and not b.ship_to_address), "transport_mode": context["mode"], "amazon_bol_ready": context["ready"], "missing_amazon_fields": context["missing"]})


def generate_bol(db: Session, ob_id, user_id, *, commit: bool = True):
    outbound = get_ob(db, ob_id)
    existing = db.scalar(
        select(BOL)
        .where(BOL.outbound_order_id == ob_id, BOL.status != BOLStatus.CANCELED)
        .order_by(BOL.id.desc())
    )
    if existing:
        from app.services.operational_document import register_generated_bol
        if register_generated_bol(db, existing, user_id) and commit:
            db.commit()
        return existing
    warehouse = outbound.warehouse
    address = ", ".join(x for x in (warehouse.address, warehouse.city, warehouse.state, warehouse.zip_code, warehouse.country) if x)
    fba = outbound.fba_shipment
    fc_address = fba.amazon_fc_address if fba else None
    b = BOL(bol_no=number(db, BOL, "BOL"), outbound_order_id=ob_id, fba_shipment_id=outbound.fba_shipment_id, customer_id=outbound.customer_id, warehouse_id=outbound.warehouse_id, carrier_id=outbound.carrier_id, ship_from_name=warehouse.warehouse_name, ship_from_address=address, ship_to_name=fc_address.fc_name if fc_address else None, ship_to_address=", ".join(x for x in (fc_address.address_line1, fc_address.city, fc_address.state, fc_address.zip_code) if x) if fc_address else None, amazon_fc_code=outbound.fc_code or (fba.amazon_fc_code if fba else None), pickup_date=to_business_date(outbound.schedule_pickup_at), appointment_time=str(outbound.delivery_appointment_time) if outbound.delivery_appointment_time else None, status=BOLStatus.GENERATED, created_by=user_id)
    db.add(b)
    db.flush()
    for allocation in outbound.allocations:
        lot = allocation.inventory_lot
        b.items.append(BOLItem(**document_identity(allocation), outbound_allocation_id=allocation.id, inventory_lot_id=lot.id, container_number=lot.container_number if fba is None else None, fc_code=lot.fc_code, marking=lot.marking, pallet_qty=allocation.allocated_pallet_qty - allocation.completed_pallet_qty, carton_qty=allocation.allocated_carton_qty - allocation.completed_carton_qty, weight_lbs=allocation.allocated_weight_lbs - allocation.completed_weight_lbs, cbm=allocation.allocated_cbm - allocation.completed_cbm, description="Outbound cargo"))
    from app.services.operational_document import register_generated_bol
    register_generated_bol(db, b, user_id)
    db.add(AuditLog(user_id=user_id, action="CREATE_BOL", entity_type="BOL", entity_id=b.id))
    if commit:
        db.commit()
    return b


def ensure_outbound_documents(db: Session, ob_id: int, user_id: int, *, commit: bool = True):
    picking = generate_picking(db, ob_id, user_id, commit=False)
    bol = generate_bol(db, ob_id, user_id, commit=False)
    if commit:
        db.commit()
    return picking, bol


def _num(value, places=2):
    return round(float(value or 0), places)


def picking_pdf(p):
    """Printable warehouse worksheet; generating it never completes a pick."""
    from xml.sax.saxutils import escape
    if 'STSong-Light' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    style = ParagraphStyle('PickingText', fontName='STSong-Light', fontSize=9, leading=13)
    title = ParagraphStyle('PickingTitle', parent=style, fontSize=19, leading=26, textColor=colors.HexColor('#17324D'))
    def cell(value):
        return Paragraph(escape(str(value if value is not None else '—')).replace('\n', '<br/>'), style)
    ob = p.outbound
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=38)
    story = [Paragraph('抓货单 / PICKING LIST', title), Spacer(1, 10)]
    for line in [f'{p.picking_no}  |  OB: {ob.ob_no}',
                 f'仓库 / Warehouse: {ob.warehouse.warehouse_name}  |  客户 / Customer: {ob.customer.customer_name if ob.customer else "—"}',
                 f'提货时间 / Pickup: {ob.schedule_pickup_at or "—"}  |  目的地 / Destination: {ob.del_code or ob.fc_code or "—"}',
                 f'状态 / Status: {PICK_NAMES.get(p.status, p.status)}',
                 '按库位顺序抓货，核对柜号和唛头后填写实抓数量；差异请记录。']:
        story.append(cell(line))
    story.append(Spacer(1, 12))
    rows = [[cell(x) for x in ['序号', '库位 / Location', '柜号 / 批次 / FC', 'BOL / PO / 唛头', '应抓\n板 / 箱', '实抓\n板 / 箱']]]
    for item in sorted(p.items, key=lambda i: (i.sequence_no or 0, i.id or 0)):
        rows.append([cell(item.sequence_no), cell(item.location.location_code if item.location else '未指定'),
                     cell('\n'.join(filter(None, [item.container_number, item.lot_no, item.fc_code]))),
                     cell('\n'.join(filter(None, [item.cargo_bol_no, f'PO: {item.po_number}' if item.po_number else None, item.marking])) or '—'), cell(f'{item.planned_pallet_qty} / {item.planned_carton_qty}'), cell('____ / ____')])
    rows.append([cell('合计'), '', '', '', cell(f'{sum(i.planned_pallet_qty for i in p.items)} / {sum(i.planned_carton_qty for i in p.items)}'), cell('____ / ____')])
    table = Table(rows, colWidths=[30, 68, 123, 123, 88, 120], repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DDEBF7')),
                              ('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#A6A6A6')),
                              ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                              ('TOPPADDING', (0, 0), (-1, -1), 9), ('BOTTOMPADDING', (0, 0), (-1, -1), 9)]))
    # Letter printable width: 612 - 60 = 552 points.
    story.extend([table, Spacer(1, 18), cell('差异 / 备注：________________________________________________________________'),
                  Spacer(1, 18), cell('抓货人：________________  复核人：________________  日期：________________'),
                  Spacer(1, 12), cell('本单仅用于抓货核对；生成或下载不会扣减库存，也不会确认出库。')])
    def footer(canvas, document):
        canvas.setFont('Helvetica', 8)
        canvas.drawString(30, 20, f'{p.picking_no} | {ob.ob_no}')
        canvas.drawRightString(582, 20, f'Page {document.page}')
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


# The warehouse's paper BOL: five cargo columns, with no extra sections.
BOL_ROWS_PER_PAGE = 20
BOL_COLUMN_WIDTHS = [156, 94, 78, 62, 130]


def _bol_form_pages(b):
    c = amazon_bol_context(b)
    cargo = []
    for item in b.items:
        allocation = getattr(item, "outbound_allocation", None)
        fba_allocation = getattr(allocation, "fba_allocation", None)
        cargo_bol = getattr(allocation, "cargo_bol", None)
        shipment = (getattr(fba_allocation, "shipment", None)
                    or getattr(cargo_bol, "fba_shipment", None))
        # Only use the parent shipment when no item-level cargo identifies another shipment.
        shipment_id = getattr(shipment, "shipment_id", None)
        reference = getattr(shipment, "reference_no", None)
        if shipment is None and cargo_bol is None:
            shipment_id = c["shipment_id"]
            reference = c["amazon_reference_id"]
        cargo.append([shipment_id or getattr(item, "cargo_bol_no", None) or item.marking or "",
                      reference or item.po_number or "", _num(item.carton_qty), "", ""])
    ship_to = _text(b.ship_to_name) or _text(b.amazon_fc_code)
    if b.amazon_fc_code and b.amazon_fc_code not in ship_to:
        ship_to += " " + b.amazon_fc_code
    header = [
        ["BILL OF LADING", "", "", "", ""],
        ["SHIP FROM", "", "BOL#", b.bol_no, ""],
        ["\n".join(filter(None, [b.ship_from_name, b.ship_from_address])), "", "", "", ""],
        ["", "", "Carrier Name:", "" if c["carrier_name"] == "To be assigned" else c["carrier_name"], ""],
        ["", "", "", "", ""],
        ["SHIP TO", "", "Trailer#", c["trailer_no"], ""],
        ["\n".join(filter(None, [ship_to, b.ship_to_address])), "", "ISA",
         "\n".join(filter(None, [c["appointment_reference"], c["appointment_time"]])), ""],
    ]
    spans = [(0, 0, 4, 0), (0, 1, 1, 1), (0, 2, 1, 4),
             (2, 1, 2, 2), (3, 1, 4, 2), (3, 3, 4, 3),
             (2, 4, 4, 4), (0, 5, 1, 5), (3, 5, 4, 5),
             (0, 6, 1, 6), (3, 6, 4, 6), (0, 27, 1, 27),
             (2, 27, 4, 27), (0, 28, 1, 28), (2, 28, 4, 28)]
    for offset in range(0, max(1, len(cargo)), BOL_ROWS_PER_PAGE):
        lines = cargo[offset:offset + BOL_ROWS_PER_PAGE]
        rows = header + lines + [["", "", "", "", ""] for _ in range(BOL_ROWS_PER_PAGE - len(lines))]
        rows += [["Shipper Signature & Date", "", "Carrier Signature & Date", "", ""], ["", "", "", "", ""]]
        yield rows, spans


def bol_xlsx(b):
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="000000")
    for page, (rows, spans) in enumerate(_bol_form_pages(b), 1):
        ws = wb.create_sheet("BOL" if page == 1 else f"BOL {page}")
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        for col, width in enumerate(BOL_COLUMN_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(col)].width = width / 7
        for row_no, row in enumerate(rows, 1):
            ws.row_dimensions[row_no].height = {1: 24, 7: 48, 29: 28}.get(row_no, 20)
            for col, value in enumerate(row, 1):
                cell = ws.cell(row_no, col, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Arial", size=16 if row_no == 1 else 10, bold=row_no in (1, 2, 6, 28))
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for x1, y1, x2, y2 in spans:
            ws.merge_cells(start_row=y1 + 1, start_column=x1 + 1, end_row=y2 + 1, end_column=x2 + 1)
        ws.print_area = "A1:E29"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def bol_pdf(b):
    from html import escape
    from reportlab.platypus import PageBreak
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter, leftMargin=40, rightMargin=40,
                            topMargin=40, bottomMargin=40, title=f"Bill of Lading {b.bol_no}")
    story = []
    for page, (rows, spans) in enumerate(_bol_form_pages(b)):
        if page:
            story.append(PageBreak())
        data = []
        for row_no, row in enumerate(rows):
            cells = []
            for value in row:
                text = f"{value:g}" if isinstance(value, float) else _text(value)
                font = "STSong-Light" if any(ord(ch) > 255 for ch in text) else ("Helvetica-Bold" if row_no in (0, 1, 5, 27) else "Helvetica")
                style = ParagraphStyle("bol-cell", fontName=font, fontSize=16 if row_no == 0 else 9,
                                       leading=18 if row_no == 0 else 11, alignment=TA_CENTER,
                                       splitLongWords=True)
                cells.append(Paragraph(escape(text).replace("\n", "<br/>") or "&#160;", style))
            data.append(cells)
        heights = [24] + [20] * 5 + [48] + [20] * 20 + [20, 28]
        table = Table(data, colWidths=BOL_COLUMN_WIDTHS, minRowHeights=heights, hAlign="CENTER")
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), .5, colors.black),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            *[("SPAN", (x1, y1), (x2, y2)) for x1, y1, x2, y2 in spans],
        ]))
        story.append(table)
    doc.build(story)
    return output.getvalue()
