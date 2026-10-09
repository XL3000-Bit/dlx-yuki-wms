"""Local documents generated exclusively from scoped, persisted business data."""
from io import BytesIO
from decimal import Decimal
from html import escape

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

PDF = 'application/pdf'
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))


def excel(sheets):
    book = Workbook()
    book.remove(book.active)
    for name, rows, widths in sheets:
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append([float(v) if isinstance(v, Decimal) else v for v in row])
        for row in sheet:
            for cell in row:
                # Treat user-entered strings as text, including Excel formula prefixes.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='203B59')
        for i, width in enumerate(widths, 1):
            sheet.column_dimensions[sheet.cell(1, i).column_letter].width = width
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = '1:1'
    target = BytesIO()
    book.save(target)
    return target.getvalue()


def pdf_document(title, blocks, *, labels=False):
    output = BytesIO()
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = 'STSong-Light'
    styles['Normal'].fontSize = 9
    styles['Normal'].leading = 13
    def p(value):
        return Paragraph(escape(str(value if value is not None else '')).replace('\n', '<br/>'), styles['Normal'])
    story = [Paragraph(escape(title), styles['Title']), Spacer(1, 12)]
    for block in blocks:
        if block is None:
            story.append(PageBreak())
        elif isinstance(block, str):
            story.extend([p(block), Spacer(1, 9)])
        else:
            rows, widths = block
            table = Table([[p(c) for c in row] for row in rows], colWidths=widths, repeatRows=1, hAlign='LEFT')
            table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#b9c4d0')),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e9eef4')),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 6),
                ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]))
            story.extend([table, Spacer(1, 12)])
    def footer(canvas, doc):
        canvas.setFont('Helvetica', 8)
        canvas.drawString(36, 22, 'LOCAL DOCUMENT | ' + str(doc.page))
    SimpleDocTemplate(output, pagesize=letter, rightMargin=36, leftMargin=36,
        topMargin=32, bottomMargin=36, title=title).build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def bol_export(row, kind):
    d = row['details']
    loads = row['loads']
    public = [
        ['BOL ID', row['bol_no'], 'Status', row['status']],
        ['OB #', row['ob_no'], 'Type', 'FBA'],
        ['Pickup Location', row['pickup'], 'Delivery', row['delivery_code']],
        ['Pickup Address', d.get('pickup_address'), 'Delivery Address', d.get('delivery_address')],
        ['Pickup Ref #', d.get('pickup_reference'), 'Delivery Ref #', d.get('delivery_reference')],
        ['OTR Carrier', row['carrier'], 'Shipping Mode', d.get('shipping_mode')],
        ['Seal #', d.get('seal_number'), 'Pro #', d.get('pro_number')],
        ['Scheduled Pickup Time', d.get('scheduled_pickup_time'), 'Actual Pickup Time', row.get('actual_pickup_time')],
        ['BOL Delivery Appointment #', d.get('delivery_appointment'), 'Delivery Appointment Time', d.get('delivery_appointment_time')],
        ['Customer Ref #', d.get('customer_reference'), 'Payment', d.get('payment')],
        ['Billing To', d.get('billing_to'), 'Delivery Time', row.get('delivery_time')],
    ]
    load_fields = [('load_id','Load ID'), ('container_number','CNTR #'), ('marking','Whs Marking'),
        ('receiver_shipment_id','Receiver Shipment ID'), ('receiver_reference_id','Receiver Reference ID'),
        ('book_qty','Book Qty'), ('actual_qty','Actual Qty'), ('current_qty','Current Qty'),
        ('weight_lbs','Weight LBS'), ('cbm','Volume CBM'), ('inbound_pallets','Inbound PLT'),
        ('whs_pallets','WHS PLT'), ('shipout_pallets','Shipout PLT'), ('remaining_pallets','Remaining PLT'),
        ('estimate_pallets','Estimate PLT'), ('markup_pallets','Markup PLT'), ('canceled_qty','Canceled Qty')]
    if kind == 'xlsx':
        history = [['Time','Operator','Request ID','Load ID','Pallets','Cartons','Weight LBS','CBM']]
        for event in row.get('workflow', {}).get('shipouts', []):
            for line in event['lines']:
                history.append([event.get('at'), event.get('operator'), event.get('request_id'), line['inventory_lot_id'],
                    *[line[k] for k in ('pallet_qty','carton_qty','weight_lbs','cbm')]])
        return excel([('BOL', [['Field','Value','Field','Value'], *public,
            ['Title Header', d.get('title_header'), 'Title Body', d.get('title_body')],
            ['Remark (Visible on BOL)', d.get('remark'), '', '']], [30,45,30,45]),
            ('Loads', [[title for _,title in load_fields], *[[x.get(key) for key,_ in load_fields] for x in loads]], [24,22,26,24,24]+[16]*12),
            ('Shipout History', history, [26,20,40,22,14,14,16,14])]), XLSX, 'xlsx'
    if kind == 'labels':
        blocks = []
        for x in loads:
            if blocks: blocks.append(None)
            blocks.extend([f"{row['bol_no']} / {x['load_id']}",
                ([['CNTR #', x['container_number']], ['Whs Marking', x['marking']],
                  ['Delivery', row['delivery_code']], ['Receiver Shipment ID', x.get('receiver_shipment_id')],
                  ['Receiver Reference ID', x.get('receiver_reference_id')], ['Book Qty', x['book_qty']],
                  ['Current Qty', x['current_qty']], ['Remaining PLT', x['remaining_pallets']]], [160,380])])
        if not blocks: raise HTTPException(409, 'Select loads before exporting labels')
        return pdf_document('Load Label', blocks), PDF, 'pdf'
    if kind != 'pdf': raise HTTPException(422, 'Supported exports: pdf, xlsx, labels')
    table = [['Receiver Shipment ID / Whs Marking','Receiver Reference ID','CTNS','PLTS','Weight LBS / CBM']]
    table += [[x.get('receiver_shipment_id') or x['marking'], x.get('receiver_reference_id'),
               x['actual_qty'], str(Decimal(str(x['whs_pallets']))+Decimal(str(x['shipout_pallets']))),
               f"{x['weight_lbs']} / {x['cbm']}"] for x in loads]
    table.append(['TOTAL', '', str(sum(Decimal(str(x['actual_qty'])) for x in loads)),
                  str(sum(Decimal(str(x['whs_pallets']))+Decimal(str(x['shipout_pallets'])) for x in loads)),
                  str(sum(Decimal(str(x['weight_lbs'])) for x in loads))])
    blocks = [d.get('title_body') or '', (public, [103,167,103,167]), (table, [172,116,55,55,142]),
        'Remark (Visible on BOL): ' + (d.get('remark') or ''),
        'Shipper Signature: __________________    Carrier Signature: __________________',
        'Date: __________________    Received By: __________________']
    return pdf_document(d.get('title_header') or 'BILL OF LADING', blocks), PDF, 'pdf'


def ocean_export(shipment, kind):
    columns = ['IB Mark','Destination','CNTR #','Book Qty','Received QTY','Inbound PLT Count','Current Qty','WHS PLT','Shipout PLT','Remaining PLT','Weight LBS','CBM','Receiver Shipment ID','Receiver Reference ID','Memo']
    rows = []
    blocks = []
    for line in shipment['data']:
        x, r, stock = line['inbound'], line['receipt'], line.get('inventory') or {}
        row = [x.get('marking'), x.get('fc_code'), x.get('container_number'), r.get('expected_qty'),
            r.get('received_qty'), r.get('inbound_pallets'), stock.get('current_qty'), stock.get('whs_pallets'),
            stock.get('shipout_pallets'), stock.get('remaining_pallets'), x.get('weight_lbs'), x.get('cbm'),
            x.get('fba_reference'), x.get('po_number'), r.get('memo')]
        rows.append(row)
        if blocks: blocks.append(None)
        blocks.extend([([['CNTR #', x.get('container_number')], ['IB Mark', x.get('marking')],
            ['Destination', x.get('fc_code')], ['BOL', line.get('cargo_bol_no')],
            ['Received QTY', r.get('received_qty')], ['Inbound PLT Count', r.get('inbound_pallets')],
            ['Receiver Shipment ID', x.get('fba_reference')], ['Receiver Reference ID', x.get('po_number')]], [160,380])])
    if kind == 'xlsx': return excel([('Ocean Inbound', [columns,*rows], [24]*len(columns))]), XLSX, 'xlsx'
    if kind == 'labels': return pdf_document('Ocean Inbound · Load Label', blocks), PDF, 'pdf'
    if kind == 'receipt':
        receipt = [['IB Mark','Destination','Received QTY','Inbound PLT','Memo']]
        receipt += [[x[0],x[1],x[4],x[5],x[-1]] for x in rows]
        return pdf_document('Ocean Inbound · Receipt', [(receipt,[130,80,80,80,170])]), PDF, 'pdf'
    raise HTTPException(422, 'Supported exports: xlsx, labels, receipt')
