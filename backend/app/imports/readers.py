import csv
import hashlib
import io
from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook


def clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def file_sha256(path: str | Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        while chunk := stream.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest()


def row_fingerprint(mapped: dict[str, Any], fields: tuple[str, ...] | None = None) -> str:
    selected = fields or tuple(sorted(mapped))
    normalized = "\x1f".join(f"{field}={str(mapped.get(field) or '').strip().casefold()}" for field in selected)
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def workbook_metadata(source: str | Path | BinaryIO) -> dict[str, Any]:
    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        names = list(workbook.sheetnames)
        estimates = {name: max(0, int(workbook[name].max_row or 0) - 1) for name in names}
        return {"sheet_names": names, "sheet_row_estimates": estimates}
    finally:
        workbook.close()


def _unique_headers(values: tuple[Any, ...]) -> list[str]:
    used: dict[str, int] = {}
    headers: list[str] = []
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value not in (None, "") else f"Column {index}"
        used[base] = used.get(base, 0) + 1
        headers.append(base if used[base] == 1 else f"{base}_{used[base]}")
    return headers


def stream_xlsx_sheet(source: str | Path | BinaryIO, sheet_name: str) -> tuple[list[str], Iterator[tuple[int, dict[str, Any]]]]:
    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"Sheet not found: {sheet_name}")
    iterator = workbook[sheet_name].iter_rows(values_only=True)
    first = next(iterator, None)
    if first is None:
        workbook.close()
        return [], iter(())
    headers = _unique_headers(first)

    def rows() -> Iterator[tuple[int, dict[str, Any]]]:
        try:
            for row_number, values in enumerate(iterator, start=2):
                row = {header: clean(values[index] if index < len(values) else None) for index, header in enumerate(headers)}
                if any(value not in (None, "") for value in row.values()):
                    yield row_number, row
        finally:
            workbook.close()

    return headers, rows()


def stream_csv(source: str | Path | BinaryIO) -> tuple[list[str], Iterator[tuple[int, dict[str, Any]]]]:
    if isinstance(source, (str, Path)):
        binary = Path(source).open("rb")
        close = True
    else:
        binary = source
        close = False
    text_stream = io.TextIOWrapper(binary, encoding="utf-8-sig", newline="")
    reader = csv.DictReader(text_stream)
    headers = [str(header).strip() for header in (reader.fieldnames or [])]

    def rows() -> Iterator[tuple[int, dict[str, Any]]]:
        try:
            for row_number, row in enumerate(reader, start=2):
                cleaned = {str(key).strip(): clean(value) for key, value in row.items()}
                if any(value not in (None, "") for value in cleaned.values()):
                    yield row_number, cleaned
        finally:
            if close:
                text_stream.close()

    return headers, rows()


def stream_tabular(source: str | Path | BinaryIO, filename: str, sheet_name: str | None = None) -> tuple[list[str], Iterator[tuple[int, dict[str, Any]]]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return stream_csv(source)
    if lower.endswith(".xlsx"):
        if not sheet_name:
            metadata = workbook_metadata(source)
            sheet_name = metadata["sheet_names"][0] if metadata["sheet_names"] else ""
            if hasattr(source, "seek"):
                source.seek(0)
        return stream_xlsx_sheet(source, sheet_name)
    raise ValueError("Only .xlsx and .csv files are supported")


def batched(rows: Iterator[tuple[int, dict[str, Any]]], batch_size: int) -> Iterator[list[tuple[int, dict[str, Any]]]]:
    batch: list[tuple[int, dict[str, Any]]] = []
    for item in rows:
        batch.append(item)
        if len(batch) >= batch_size:
            yield batch
            batch = []
    if batch:
        yield batch


def read_tabular(content: bytes, filename: str, sheet_name: str | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    stream = io.BytesIO(content)
    headers, iterator = stream_tabular(stream, filename, sheet_name)
    return headers, [row for _, row in iterator]
