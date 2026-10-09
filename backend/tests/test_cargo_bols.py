from tests.dispatch_history import mark_historical_dispatched
from datetime import date
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.security import create_access_token
from app.models import CargoBOL, Customer, InventoryLot, OutboundInventoryAllocation, OutboundOrder, PickingList, PickingStatus
from app.models.user import ScopeMode


def stock(client, seed, name, customer_id=None, receive=True):
    response = client.post('/api/v1/inbound', json={
        'container_number': name, 'po_number': name, 'customer_id': customer_id or seed['customer'].id,
        'warehouse_id': seed['warehouse'].id, 'received_date': str(date.today()), 'fc_code': 'ONT8',
        'pallet_qty': 10, 'carton_qty': 20, 'weight_lbs': 1000, 'cbm': 5,
        'location_id': seed['location'].id, 'status': 3,
    })
    assert response.status_code == 201, response.text
    inbound = response.json()
    lot = None
    if receive:
        response = client.post(f"/api/v1/inbound/{inbound['id']}/receive-to-inventory")
        assert response.status_code == 200, response.text
        lot = response.json()
    rows = client.get('/api/v1/cargo-bols', params={'q': name}).json()['data']
    cargo = next(row for row in rows if row['inbound_id'] == inbound['id'])
    return cargo, lot


def order(seed, **changes):
    return {'warehouse_id': seed['warehouse'].id, 'customer_id': seed['customer'].id,
            'carrier_id': seed['carrier'].id, 'ob_type': 'STANDARD', 'delivery_type': 'FTL', **changes}


def create_from(client, seed, ids, **changes):
    return client.post('/api/v1/cargo-bols/create-outbound', json={'outbound': order(seed, **changes), 'bol_ids': ids})


def test_filter_pagination_export_and_literal_references(client, db, seed):
    first, _ = stock(client, seed, 'PO_%_MATCH')
    stock(client, seed, 'PO_ABC_MATCH')
    formula, _ = stock(client, seed, '=1+1')
    filters = {'warehouse_id': seed['warehouse'].id, 'customer_id': seed['customer'].id,
               'source_type': 'INBOUND', 'del_code': ' ont8 ', 'remaining_only': True}
    listing = client.get('/api/v1/cargo-bols', params={**filters, 'page_size': 1}).json()
    assert len(listing['data']) == 1 and listing['meta']['total'] == 3
    assert {key: float(value) for key, value in listing['meta']['totals'].items()} == {
        'available_pallet_qty': 30, 'available_carton_qty': 60,
        'available_weight_lbs': 3000, 'available_cbm': 15,
    }
    filtered = client.get('/api/v1/cargo-bols', params={'q': '%_'}).json()
    assert float(filtered['meta']['totals']['available_pallet_qty']) == 10
    empty = client.get('/api/v1/cargo-bols', params={**filters, 'warehouse_id': 99999}).json()
    assert all(float(value) == 0 for value in empty['meta']['totals'].values())
    assert client.get('/api/v1/cargo-bols', params={'q': '%_'}).json()['data'][0]['id'] == first['id']
    assert client.get('/api/v1/cargo-bols', params={'q': '%_'}).json()['meta']['total'] == 1
    assert client.get('/api/v1/cargo-bols', params={'q': formula['bol_no']}).json()['meta']['total'] == 1
    assert client.get('/api/v1/cargo-bols', params={**filters, 'source_type': 'FBA'}).json()['meta']['total'] == 0
    assert client.get('/api/v1/cargo-bols', params={**filters, 'warehouse_id': 99999}).json()['meta']['total'] == 0
    assert client.get('/api/v1/cargo-bols', params={**filters, 'del_code': 'ONT'}).json()['meta']['total'] == 0
    export = client.get('/api/v1/cargo-bols/export.xlsx', params=filters)
    assert export.status_code == 200, export.text
    sheet = load_workbook(BytesIO(export.content)).active
    assert sheet.max_row == 4 and sheet.freeze_panes == 'A2'
    reference = next(cell for row in sheet for cell in row if cell.value == '=1+1')
    assert reference.data_type == 's'


def test_create_and_assign_are_atomic_on_stale_inventory(client, db, seed):
    available, lot = stock(client, seed, 'AVAILABLE')
    exhausted, _ = stock(client, seed, 'EXHAUSTED')
    assert create_from(client, seed, [exhausted['id']]).status_code == 201
    count = db.scalar(select(func.count()).select_from(OutboundOrder))
    failed = create_from(client, seed, [available['id'], exhausted['id']])
    assert failed.status_code == 409, failed.text
    db.expire_all()
    assert db.scalar(select(func.count()).select_from(OutboundOrder)) == count
    assert db.get(InventoryLot, lot['id']).available_pallet_qty == 10
    existing = client.post('/api/v1/outbounds', json=order(seed)).json()
    failed = client.post('/api/v1/cargo-bols/assign', json={'outbound_id': existing['id'], 'bol_ids': [available['id'], exhausted['id']]})
    assert failed.status_code == 409, failed.text
    db.expire_all()
    assert db.get(InventoryLot, lot['id']).available_carton_qty == 20
    assert db.scalar(select(OutboundInventoryAllocation).where(OutboundInventoryAllocation.outbound_order_id == existing['id'])) is None


def test_duplicate_selection_allocates_once_and_cannot_allocate_again(client, db, seed):
    cargo, lot = stock(client, seed, 'DUPLICATE')
    created = create_from(client, seed, [cargo['id'], cargo['id']])
    assert created.status_code == 201, created.text
    result = created.json()
    assert result['bol_count'] == 1 and len(result['allocation_ids']) == 1
    detail = client.get(f"/api/v1/outbounds/{result['ob_id']}/workbench-detail").json()
    assert detail['basic']['customer_id'] == seed['customer'].id
    assert detail['allocations'][0]['cargo_bol_no'] == cargo['bol_no']
    assert float(detail['basic']['allocated_weight_lbs']) == 1000
    assert client.post('/api/v1/cargo-bols/assign', json={'outbound_id': result['ob_id'], 'bol_ids': [cargo['id']]}).status_code == 409
    db.expire_all()
    assert db.get(InventoryLot, lot['id']).available_pallet_qty == 0
    assert db.get(InventoryLot, lot['id']).allocated_pallet_qty == 10
    assert client.get('/api/v1/cargo-bols', params={'remaining_only': True}).json()['meta']['total'] == 0


def test_scope_customer_mismatch_and_viewer_write_denial(client, db, seed):
    own, _ = stock(client, seed, 'OWN-CARGO')
    other = Customer(customer_code='OTHER', customer_name='Other customer')
    db.add(other); db.commit()
    foreign, _ = stock(client, seed, 'FOREIGN-CARGO', customer_id=other.id)
    assert create_from(client, seed, [own['id'], foreign['id']]).status_code == 409
    created = client.post('/api/v1/outbounds', json=order(seed)).json()
    detail = client.get(f"/api/v1/outbounds/{created['id']}/workbench-detail").json()
    assert len(detail['remaining_sources']) == 1
    viewer = seed['viewer']
    viewer.customer_scope_mode = ScopeMode.SELECTED
    viewer.customers = [seed['customer']]
    db.commit()
    client.headers['Authorization'] = f'Bearer {create_access_token(str(viewer.id))}'
    assert client.get('/api/v1/cargo-bols').json()['meta']['total'] == 1
    assert float(client.get('/api/v1/cargo-bols').json()['meta']['totals']['available_pallet_qty']) == 10
    assert client.get(f"/api/v1/cargo-bols/{foreign['id']}").status_code == 404
    export = client.get('/api/v1/cargo-bols/export.xlsx')
    assert load_workbook(BytesIO(export.content)).active.max_row == 2
    assert create_from(client, seed, [own['id']]).status_code == 403
    assert client.post('/api/v1/cargo-bols/assign', json={'outbound_id': created['id'], 'bol_ids': [own['id']]}).status_code == 403


def test_unreceived_cargo_cannot_create_outbound(client, db, seed):
    cargo, _ = stock(client, seed, 'NOT-RECEIVED', receive=False)
    assert cargo['can_allocate'] is False
    response = create_from(client, seed, [cargo['id']])
    assert response.status_code == 409, response.text
    assert db.scalar(select(func.count()).select_from(OutboundOrder)) == 0
    assert create_from(client, seed, []).status_code == 422
    assert create_from(client, seed, [cargo['id']], customer_id=None).status_code == 422


def test_fba_multiple_lots_remaining_and_completion_do_not_double_deduct(client, db, seed):
    _, lot1 = stock(client, seed, 'FBA-LOT-1')
    _, lot2 = stock(client, seed, 'FBA-LOT-2')
    fba = client.post('/api/v1/fba', json={'warehouse_id': seed['warehouse'].id,
        'customer_id': seed['customer'].id, 'amazon_fc_code': 'ONT8'}).json()
    reservations = []
    for lot in (lot1, lot2):
        response = client.post(f"/api/v1/fba/{fba['id']}/allocate", json={'inventory_lot_id': lot['id'],
            'pallet_qty': 10, 'carton_qty': 20, 'weight_lbs': 1000, 'cbm': 5})
        assert response.status_code == 200, response.text
        reservations.append(response.json())
    cargo = client.get('/api/v1/cargo-bols', params={'source_type': 'FBA'}).json()['data'][0]
    assert float(cargo['available_pallet_qty']) == 20
    assert create_from(client, seed, [cargo['id']]).status_code == 409
    first = client.post('/api/v1/outbounds', json=order(seed, ob_type='FBA', fba_shipment_id=fba['id'])).json()
    response = client.post(f"/api/v1/outbounds/{first['id']}/allocate", json={'inventory_lot_id': lot1['id'],
        'fba_allocation_id': reservations[0]['id'], 'pallet_qty': 4, 'carton_qty': 8, 'weight_lbs': 400, 'cbm': 2})
    assert response.status_code == 200, response.text
    assert float(client.get(f"/api/v1/cargo-bols/{cargo['id']}").json()['available_pallet_qty']) == 16
    created = create_from(client, seed, [cargo['id']], ob_type='FBA', fba_shipment_id=fba['id'])
    assert created.status_code == 201, created.text
    result = created.json()
    assert len(result['allocation_ids']) == 2
    detail = client.get(f"/api/v1/cargo-bols/{cargo['id']}").json()
    assert float(detail['available_pallet_qty']) == 0
    links = [row for row in detail['outbounds'] if row['id'] == result['ob_id']]
    assert len(links) == 2 and len({row['allocation_id'] for row in links}) == 2
    db.expire_all()
    assert db.get(InventoryLot, lot1['id']).allocated_pallet_qty == 10
    assert db.get(InventoryLot, lot2['id']).allocated_pallet_qty == 10
    assert client.post(f"/api/v1/outbounds/{first['id']}/confirm").status_code == 200
    picking = db.scalar(select(PickingList).where(PickingList.outbound_order_id == first['id']))
    for item in picking.items:
        item.picked_pallet_qty = item.planned_pallet_qty
    picking.status = PickingStatus.COMPLETED
    db.commit()
    assert client.post(f"/api/v1/outbounds/{first['id']}/dispatch").status_code == 409
    mark_historical_dispatched(db, first['id'])
    assert client.post(f"/api/v1/outbounds/{first['id']}/complete").status_code == 200
    assert float(client.get(f"/api/v1/cargo-bols/{cargo['id']}").json()['available_pallet_qty']) == 0
    assert client.post(f"/api/v1/outbounds/{result['ob_id']}/cancel").status_code == 200
    detail = client.get(f"/api/v1/cargo-bols/{cargo['id']}").json()
    assert float(detail['available_pallet_qty']) == 16
    assert float(detail['available_carton_qty']) == 32
    db.expire_all()
    assert db.get(InventoryLot, lot1['id']).allocated_pallet_qty == 6
    assert db.get(InventoryLot, lot1['id']).available_pallet_qty == 0


def test_cancel_is_not_reported_as_dispatch_completion(client, seed):
    created = client.post('/api/v1/outbounds', json=order(seed)).json()
    assert client.post(f"/api/v1/outbounds/{created['id']}/cancel").status_code == 200
    detail = client.get(f"/api/v1/outbounds/{created['id']}/workbench-detail").json()
    assert next(step for step in detail['workflow'] if step['key'] == 'dispatch')['status'] == 'pending'


def test_migration_backfills_existing_cargo_without_changing_stock(client, db, seed, monkeypatch):
    import importlib.util
    from pathlib import Path
    from sqlalchemy import delete, update
    from app.models import FBAInventoryAllocation, ImportJob, InboundRecord
    from app.models.import_job import ImportModule

    kept, _ = stock(client, seed, 'MIGRATION-KEEP')
    inbound_cargo, _ = stock(client, seed, 'MIGRATION-INBOUND')
    standard = create_from(client, seed, [inbound_cargo['id']]).json()
    _, fba_lot = stock(client, seed, 'MIGRATION-FBA')
    fba = client.post('/api/v1/fba', json={'warehouse_id': seed['warehouse'].id,
        'customer_id': seed['customer'].id, 'amazon_fc_code': 'ONT8'}).json()
    response = client.post(f"/api/v1/fba/{fba['id']}/allocate", json={'inventory_lot_id': fba_lot['id'],
        'pallet_qty': 10, 'carton_qty': 20, 'weight_lbs': 1000, 'cbm': 5})
    assert response.status_code == 200, response.text
    fba_cargo = client.get('/api/v1/cargo-bols', params={'source_type': 'FBA'}).json()['data'][0]
    fba_ob = create_from(client, seed, [fba_cargo['id']], ob_type='FBA', fba_shipment_id=fba['id']).json()
    historic = client.post('/api/v1/outbounds', json=order(seed)).json()
    job = ImportJob(module=ImportModule.OUTBOUND, file_name='history.csv', original_file_name='history.csv',
                    profile_code='WEST_COAST_4_0_HISTORY', created_by=seed['admin'].id)
    db.add(job)
    db.flush()
    db.get(OutboundOrder, historic['id']).import_job_id = job.id
    db.flush()
    # Simulate records created before cargo identities were introduced.
    db.execute(update(OutboundInventoryAllocation).values(cargo_bol_id=None))
    db.execute(delete(CargoBOL).where(CargoBOL.id != kept['id']))
    db.flush()

    def quantities():
        return [list(db.execute(select(*[column for column in model.__table__.c
                     if column.name != 'cargo_bol_id'])).all())
                for model in (InventoryLot, FBAInventoryAllocation, OutboundInventoryAllocation)]

    before = quantities()
    spec = importlib.util.spec_from_file_location('cargo_migration', Path(__file__).parents[1] /
        'alembic/versions/20260909_0033_backfill_cargo_bols.py')
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    monkeypatch.setattr(migration.op, 'get_bind', lambda: db.connection())
    migration.upgrade()
    identities = list(db.execute(select(CargoBOL.__table__).order_by(CargoBOL.id)).all())
    migration.upgrade()
    assert identities == list(db.execute(select(CargoBOL.__table__).order_by(CargoBOL.id)).all())
    assert quantities() == before
    db.expire_all()
    assert db.get(CargoBOL, kept['id']).inbound_id == kept['inbound_id']
    assert db.scalar(select(func.count()).select_from(CargoBOL).where(CargoBOL.inbound_id.is_not(None))) == db.scalar(select(func.count()).select_from(InboundRecord))
    assert db.scalar(select(CargoBOL).where(CargoBOL.history_outbound_id == historic['id'])) is not None
    assert db.scalar(select(func.count()).select_from(CargoBOL).where(CargoBOL.history_outbound_id.is_not(None))) == 1
    standard_link = db.scalar(select(OutboundInventoryAllocation).where(OutboundInventoryAllocation.outbound_order_id == standard['ob_id']))
    fba_link = db.scalar(select(OutboundInventoryAllocation).where(OutboundInventoryAllocation.outbound_order_id == fba_ob['ob_id']))
    assert db.get(CargoBOL, standard_link.cargo_bol_id).inbound_id == inbound_cargo['inbound_id']
    assert db.get(CargoBOL, fba_link.cargo_bol_id).fba_shipment_id == fba['id']
