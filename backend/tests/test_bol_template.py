from io import BytesIO
from types import SimpleNamespace as NS

import pytest
from openpyxl import load_workbook

from app.services import picking_bol as service


def sample_bol(monkeypatch, count=4):
    monkeypatch.setattr(service, "amazon_bol_context", lambda b: {
        "shipment_id": "PARENT", "amazon_reference_id": "PARENT-REF",
        "carrier_name": "To be assigned", "seal_no": "634533",
        "trailer_no": "", "appointment_reference": "93426047990",
        "appointment_time": "09/12/2026 22:00 PDT",
    })
    return NS(bol_no="SAMPLE-090126", ship_from_name="DLX Logistics International LLC",
              ship_from_address="4450 EDISON AVE, CHINO 91710", ship_to_name="TCY1",
              ship_to_address="2690 Arch Airport Rd, Stockton, CA 95206", amazon_fc_code="TCY1",
              items=[NS(outbound_allocation=NS(fba_allocation=NS(shipment=NS(
                  shipment_id=f"FBA-{i}", reference_no=f"REF-{i}")), cargo_bol=None),
                  container_number="LEGACY-CONTAINER", cargo_bol_no=f"BOL-{i}", marking="", po_number="", carton_qty=i+1)
                     for i in range(count)])


def test_template_uses_each_shipment_and_cartons(monkeypatch):
    bol = sample_bol(monkeypatch)
    bol.items[1].outbound_allocation = NS(fba_allocation=None, cargo_bol=NS(fba_shipment=None))
    rows, _ = next(service._bol_form_pages(bol))
    assert rows[7] == ["FBA-0", "REF-0", 1, "", ""]
    assert rows[8] == ["BOL-1", "", 2, "", ""]
    assert rows[27][0] == "Shipper Signature & Date"


def test_template_paginates_and_preserves_literal_strings(monkeypatch):
    bol = sample_bol(monkeypatch, 21)
    bol.bol_no = "=1+1"
    wb = load_workbook(BytesIO(service.bol_xlsx(bol)))
    assert wb.sheetnames == ["BOL", "BOL 2"]
    assert wb["BOL"]["D2"].data_type == "s"
    assert wb["BOL 2"]["A8"].value == "FBA-20"
    assert wb["BOL 2"]["A9"].value is None
    assert wb["BOL 2"]["C28"].value == "Carrier Signature & Date"
    assert service.bol_pdf(bol).startswith(b"%PDF-")


@pytest.mark.parametrize("field", ["container", "seal"])
def test_bol_no_container_or_seal_field(monkeypatch, field):
    bol = sample_bol(monkeypatch, 21)
    forbidden = [field, "LEGACY-CONTAINER" if field == "container" else "634533"]
    wb = load_workbook(BytesIO(service.bol_xlsx(bol)))
    text = " ".join(str(cell.value or "") for ws in wb for row in ws for cell in row)
    assert all(word.lower() not in text.lower() for word in forbidden)

    # Disable compression only in this test so we can inspect actual PDF text streams.
    document = service.SimpleDocTemplate
    monkeypatch.setattr(service, "SimpleDocTemplate", lambda *args, **kw:
                        document(*args, **kw, pageCompression=0))
    pdf = service.bol_pdf(bol)
    assert pdf.startswith(b"%PDF-")
    assert b"BILL OF LADING" in pdf
    assert all(word.lower().encode() not in pdf.lower() for word in forbidden)


@pytest.mark.parametrize("seal", [None, "LEGACY-SEAL"])
def test_bol_does_not_require_or_read_seal(seal):
    class Load:
        @property
        def seal_no(self):
            raise AssertionError("BOL must not read load seal")
    outbound = NS(delivery_type="FTL", truck_type=None, ob_type="FBA",
                  fba_shipment=NS(reference_no="REF1", shipment_id="FBA1"),
                  bol_reference="PRO1", load=Load())
    bol = NS(outbound=outbound, ship_to_address="Destination", amazon_fc_code="ONT8",
             seal_no=seal, container_number="LEGACY-CONTAINER")
    payload = service.amazon_bol_context(bol)
    assert payload["mode"] == "FTL"
    assert payload["ready"] is True
    assert payload["missing"] == []
    assert not any("container" in key.lower() for key in payload)
    assert not any("seal" in key.lower() for key in payload)
