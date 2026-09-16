from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace as NS

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.services import picking_bol as service


def sample_bol(kind="fba_type", seal="HIST-SEAL-987"):
    fba = NS(reference_no="AMZ-REF", shipment_id="SHIP-123", st_number="ST-456", po_number="PO-789")
    outbound = NS(ob_type="FBA" if kind == "fba_type" else "OTHER",
                  fba_shipment=fba if kind == "fba_outbound_relation" else None,
                  load=NS(seal_no=seal, tractor_no="TRACTOR-1", trailer_no="TRAILER-2"),
                  delivery_type="FTL", ob_no="OB-SAMPLE", bol_reference="PRO-123", po_number="PO-789")
    return NS(outbound=outbound, fba_shipment=fba if kind == "fba_bol_relation" else None,
              carrier=NS(carrier_name="Sample Carrier", scac="TEST"),
              bol_no="BOL-SAMPLE", ship_from_name="Sample Warehouse", ship_from_address="1 Sample Way",
              ship_to_name="Sample FC", ship_to_address="2 Example Road", amazon_fc_code="ONT8",
              pickup_date="2026-09-16", items=[NS(container_number="HIST-CNTR-123", fc_code="ONT8",
              marking="MARK-1", description="Sample merchandise", pallet_qty=Decimal(2), carton_qty=Decimal(20),
              weight_lbs=Decimal(100), cbm=Decimal(1))])


def export_text(bol, format):
    if format == "xlsx":
        ws = load_workbook(BytesIO(service.bol_xlsx(bol))).active
        text = "\n".join(str(cell.value) for row in ws for cell in row if cell.value is not None)
    else:
        # Read the completed PDF, including page callbacks; do not intercept build.
        pdf = PdfReader(BytesIO(service.bol_pdf(bol)))
        assert len(pdf.pages) > 0
        pages = [page.extract_text() for page in pdf.pages]
        assert all(page.strip() for page in pages)
        text = "\n".join(pages)
    # Wrapping differs between PDF and Excel, but must not change content.
    return " ".join(text.split())


@pytest.mark.parametrize("format", ["pdf", "xlsx"])
@pytest.mark.parametrize("kind", ["fba_type", "fba_bol_relation", "fba_outbound_relation", "other"])
@pytest.mark.parametrize("seal", ["HIST-SEAL-987", ""])
def test_export_scope_and_readiness(format, kind, seal):
    bol = sample_bol(kind, seal)
    assert (bool(bol.fba_shipment), bool(bol.outbound.fba_shipment), bol.outbound.ob_type == "FBA") == {
        "fba_bol_relation": (True, False, False),
        "fba_outbound_relation": (False, True, False),
        "fba_type": (False, False, True),
        "other": (False, False, False),
    }[kind]
    before = service.amazon_bol_context(bol)
    text = export_text(bol, format)
    assert service.amazon_bol_context(bol) == before
    assert bol.items[0].container_number == "HIST-CNTR-123"
    assert bol.outbound.load.seal_no == seal
    for retained in ["BOL-SAMPLE", "OB-SAMPLE", "PRO-123", "TRACTOR-1", "TRAILER-2", "MARK-1", "Sample merchandise", "PO-789",
                     "Sample Warehouse", "1 Sample Way", "Sample FC", "2 Example Road", "Sample Carrier", "ONT8",
                     "Non-partner carriers must schedule a delivery appointment.",
                     "Keep Amazon box and pallet labels visible."]:
        assert retained in text
    if kind != "other":
        for removed in ["Container", "Trailer Seal (FTL)", "HIST-CNTR-123", "HIST-SEAL-987",
                        "Trailer seal number", "FTL/intermodal loads require a trailer seal."]:
            assert removed not in text
        assert "TO BE PROVIDED" not in text
        if not seal:
            assert "Trailer seal number" in before["missing"]
            assert not before["ready"]
            assert "PRE-DISPATCH CHECK" in text
        if kind in ("fba_bol_relation", "fba_outbound_relation"):
            for reference in ("AMZ-REF", "SHIP-123", "ST-456"):
                assert reference in text
            if not seal:
                assert before["missing"] == ["Trailer seal number"]
        if kind == "fba_type":
            assert "Amazon Reference ID; FBA Shipment ID" in text
    else:
        assert "Container" in text and "HIST-CNTR-123" in text
        assert "Trailer Seal (FTL)" in text
        assert "FTL/intermodal loads require a trailer seal." in text
        assert (seal or "TO BE PROVIDED") in text


def test_fba_snapshot_and_picking_container_preserved(client, seed, db):
    from app.models import BOL, OutboundOrder
    from test_picking_bol import make_ob
    outbound, _ = make_ob(client, seed, "-FBA-EXPORT")
    db.get(OutboundOrder, outbound["id"]).ob_type = "FBA"
    db.commit()
    docs = client.post(f"/api/v1/outbounds/{outbound['id']}/documents/ensure").json()
    bol = db.get(BOL, docs["bol_id"])
    assert bol.items[0].container_number == "PK-CNTR-FBA-EXPORT"
    for format in ("pdf", "xlsx"):
        assert client.get(f"/api/v1/bols/{bol.id}/{format}").status_code == 200
    assert bol.items[0].container_number == "PK-CNTR-FBA-EXPORT"
    response = client.get(f"/api/v1/picking-lists/{docs['picking_list_id']}/xlsx")
    assert response.status_code == 200
    ws = load_workbook(BytesIO(response.content)).active
    assert ws["E1"].value == "Container"
    assert ws["E2"].value == "PK-CNTR-FBA-EXPORT"
