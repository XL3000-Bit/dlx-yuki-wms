
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import select, func
from app.models import ImportJob, ImportRow, OutboundOrder, InventoryLot
from app.models.import_job import ImportModule
from app.services.convert_west_history import PROFILE
from app.core.security import create_access_token

def test_historical_bol_metadata_and_transfer_code(client, db, seed):
    job = ImportJob(module=ImportModule.OUTBOUND, file_name='history.xlsx',
        original_file_name='history.xlsx', profile_code=PROFILE, source_sheet='出库',
        created_by=seed['admin'].id)
    db.add(job); db.flush()
    order = OutboundOrder(ob_no='HOB-TEST', warehouse_id=seed['warehouse'].id,
        created_by=seed['admin'].id, status=1, ob_type='FBA', import_job_id=job.id, fc_code='SCK1')
    db.add(order); db.flush()
    db.add(ImportRow(import_job_id=job.id, row_number=2,
        raw_data={'BOL': 1, '重量(lb)': 33598.312},
        created_entity_type='OUTBOUND', created_entity_id=order.id))
    db.commit()
    response = client.get('/api/v1/outbounds/workbench')
    assert response.status_code == 200, response.text
    row = response.json()['data'][0]
    assert row['bol_no'] is None and row['bol_type'] is None
    assert float(row['weight_lbs']) == 33598.312
    assert row['del_code'] == 'SCK1'
    endpoint = '/api/v1/outbounds/workbench/transfer-code'
    invalid = client.patch(endpoint, json={'ids': [order.id, 999999], 'transfer_code': 'BAD'})
    assert invalid.status_code == 404
    db.refresh(order)
    assert order.transfer_code is None
    result = client.patch(endpoint, json={'ids': [order.id], 'transfer_code': '=1+1'})
    assert result.status_code == 200, result.text
    db.refresh(order)
    assert order.transfer_code == '=1+1' and order.status == 1
    assert db.scalar(select(func.count()).select_from(InventoryLot)) == 0
    assert len(client.get('/api/v1/outbounds/workbench', params={'q': '=1+1'}).json()['data']) == 1
    export = client.post('/api/v1/outbounds/workbench/bol-export', json={'ids': [order.id]})
    assert export.status_code == 200, export.text
    ws = load_workbook(BytesIO(export.content)).active
    assert ws['H2'].value == '=1+1' and ws['H2'].data_type == 's'
    assert ws['I2'].value == 33598.312
    client.headers['Authorization'] = f"Bearer {create_access_token(str(seed['viewer'].id))}"
    assert client.patch(endpoint, json={'ids': [order.id], 'transfer_code': 'DENIED'}).status_code == 403
