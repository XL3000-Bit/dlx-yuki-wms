from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "DLX_Yuki_WMS_员工操作手册_零基础版.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "18324A"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "E8F4EC"
PALE_YELLOW = "FFF4CC"
PALE_RED = "FDE9E7"
GRAY = "F3F5F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"
RED = "B42318"
GREEN = "18794E"
ORANGE = "B54708"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_grid(table, widths):
    """Force Word to honor the intended fixed column widths."""
    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def set_repeat_header(section, text):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MID_GRAY)

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DLX Yuki WMS · 员工操作手册   |   ")
    r.font.name = "Calibri"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MID_GRAY)
    add_page_field(p)


def set_run_font(run, size=None, bold=None, color=None, east_asia="Microsoft YaHei"):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_keep(paragraph, next_=False):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_lines = OxmlElement("w:keepLines")
    p_pr.append(keep_lines)
    if next_:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, color, before, after in (
        ("Title", 30, NAVY, 0, 10),
        ("Subtitle", 13, MID_GRAY, 0, 12),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Step Title" not in doc.styles:
        style = doc.styles.add_style("Step Title", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = True


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.375 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def add_callout(doc, title, text, kind="info"):
    palette = {
        "info": (PALE_BLUE, DARK_BLUE),
        "success": (PALE_GREEN, GREEN),
        "warn": (PALE_YELLOW, ORANGE),
        "danger": (PALE_RED, RED),
    }
    fill, color = palette[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid(table, [9360])
    set_cell_width(table.cell(0, 0), 9360)
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell_margins(cell, 140, 180, 140, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, 11, True, color)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)
    keep_row(table.rows[0])
    return table


def add_steps(doc, steps):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid(table, [850, 8510])
    for idx, (action, check) in enumerate(steps, 1):
        row = table.add_row()
        keep_row(row)
        set_cell_width(row.cells[0], 850)
        set_cell_width(row.cells[1], 8510)
        shade(row.cells[0], BLUE)
        shade(row.cells[1], WHITE if idx % 2 else "F8FAFC")
        for cell in row.cells:
            cell_margins(cell, 80, 130, 80, 130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = row.cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(idx))
        set_run_font(r, 12, True, WHITE)
        p = row.cells[1].paragraphs[0]
        r = p.add_run(action)
        set_run_font(r, 11, True, NAVY)
        if check:
            p2 = row.cells[1].add_paragraph()
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run("成功标志：" + check)
            set_run_font(r2, 9, False, GREEN)
    return table


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        set_table_grid(table, widths)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, PALE_BLUE)
        cell_margins(cell)
        if widths:
            set_cell_width(cell, widths[i])
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_run_font(r, 9, True, NAVY)
    set_repeat_table_header(table.rows[0])
    keep_row(table.rows[0])
    for ridx, values in enumerate(rows):
        row = table.add_row()
        keep_row(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            shade(cell, WHITE if ridx % 2 == 0 else "F8FAFC")
            cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[i])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, 9)
    return table


def page_break(doc):
    # Section headings carry page_break_before. This avoids Word pushing a
    # stand-alone break paragraph onto a new page and creating a blank page.
    return None


def section_title(doc, n, title, subtitle=None, break_before=True):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = break_before
    p.add_run(f"{n}  {title}")
    if subtitle:
        p2 = doc.add_paragraph(subtitle)
        p2.style = doc.styles["Subtitle"]
        p2.paragraph_format.space_after = Pt(10)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    set_repeat_header(section, "员工操作手册 · 零基础版")

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(48)
    r = p.add_run("DLX")
    set_run_font(r, 18, True, BLUE)
    r = p.add_run("  YUKI WMS")
    set_run_font(r, 18, True, NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("员工操作手册")
    set_run_font(r, 30, True, NAVY)
    p = doc.add_paragraph()
    r = p.add_run("零基础版 · 一线仓库员工适用")
    set_run_font(r, 16, True, BLUE)
    p.paragraph_format.space_after = Pt(28)

    add_callout(doc, "这本手册怎么用", "先看“安全红线”，再按左侧菜单找到对应章节。每个流程都写了：点哪里、填什么、看到什么算成功、出错怎么办。", "info")
    doc.add_paragraph()
    add_table(doc, ["适用工作", "包含内容"], [
        ("收货与入库", "Inbound、Receive to Inventory"),
        ("库存与柜况", "Inventory、Container Tracking"),
        ("拣货与出库", "Picking List、PDA、Outbound、BOL、Loads"),
        ("协作与留档", "Work Orders、Trouble Shoot、Documents & POD"),
    ], [2500, 6860])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    r = p.add_run("版本 1.0  |  2026-09-02")
    set_run_font(r, 10, False, MID_GRAY)
    p = doc.add_paragraph("系统按钮为英文，解释为中文。实际能看到的菜单取决于账号权限。")
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        set_run_font(run, 9, False, MID_GRAY)

    page_break(doc)
    section_title(doc, "先看", "6条安全红线", "新人最重要的一页：不确定时，先停手，不要靠猜。")
    rules = [
        ("账号只给自己用", "不要把密码发给同事，也不要让浏览器记住公共电脑上的密码。离开工位要退出。"),
        ("带动作的按钮会改数据", "Save、Create、Complete、Dispatch、Adjust、Archive、Cancel 都可能改变业务记录。点前再核对单号和数量。"),
        ("实物没完成，不点完成", "仓库现场尚未收完、拣完、装完或发走时，不要提前点 Complete / Dispatch。"),
        ("一次只点一下", "保存后等待绿色成功提示或页面刷新。不要连续双击，避免重复创建。"),
        ("离线保存不等于已上传", "PDA 显示“保存到本机”时，数据还没到服务器。恢复网络后必须等到同步成功。"),
        ("错了就留证据", "记下业务单号、时间和你刚做的操作；截图错误提示；找主管。不要反复调整把问题越改越乱。"),
    ]
    add_steps(doc, rules)
    add_callout(doc, "必须马上停手的情况", "单号对不上、实物数量对不上、库位标签看不清、出现红色错误、系统卡住但不确定是否保存、离线记录无法同步。", "danger")

    page_break(doc)
    section_title(doc, "0", "30秒认识系统", "你只需要记住：左边找功能，上面找搜索，右上角找账号。")
    add_table(doc, ["位置", "你会看到", "用来做什么"], [
        ("左侧菜单", "InBound / Warehouse / OutBound / Operations", "进入收货、库存、出库和异常页面"),
        ("页面上方", "Search / Reset / Refresh / Export", "查找、清空筛选、刷新、导出"),
        ("表格里的蓝色单号", "如 OB#、Exception No", "点击打开详情"),
        ("表格右侧", "Actions / More / …", "打开该行可执行的操作"),
        ("右上角头像", "账号菜单 / Sign out", "查看账号、退出系统"),
    ], [1600, 3000, 4760])
    doc.add_heading("最常见按钮翻译", level=2)
    add_table(doc, ["英文", "中文", "新手提示"], [
        ("Search", "搜索", "先输入条件，再点它"),
        ("Reset", "重置", "清空筛选，不删除数据"),
        ("Refresh", "刷新", "重新读取服务器数据"),
        ("Create", "新建", "会产生新业务记录"),
        ("Save", "保存", "只点一次，等提示"),
        ("Complete", "完成", "实物完成后才能点"),
        ("Cancel", "取消/作废", "弹窗里的 Cancel 常是关闭；业务 Actions 里的 Cancel 是作废"),
        ("Export / Download", "导出/下载", "文件通常在电脑“下载”文件夹"),
    ], [1600, 1700, 6060])
    add_callout(doc, "快速判断是否保存成功", "优先看右上角是否出现绿色提示，再看状态或列表是否更新。如果没有明确成功提示，不要重复点；先 Refresh，再按单号搜索确认。", "success")

    page_break(doc)
    section_title(doc, "1", "登录、退出与每日检查")
    doc.add_heading("登录", level=2)
    add_steps(doc, [
        ("打开公司提供的 WMS 网址。", "看到 DLX Yuki WMS 登录页"),
        ("在 USERNAME 输入用户名，在 PASSWORD 输入密码。", "密码输入后会显示为圆点"),
        ("点击 INITIALIZE SESSION。", "进入 Dashboard，左侧出现菜单"),
    ])
    add_callout(doc, "登录失败", "看到 “Invalid username/password or backend unavailable.”：先检查大小写和网络；只重试一次。仍失败就把错误截图发给主管或系统管理员。", "warn")
    doc.add_heading("上班前 1 分钟", level=2)
    for s in ["确认当前账号是你本人。", "确认网络正常，页面顶部没有离线提示。", "进入今天负责的模块，点 Refresh。", "确认仓库 Warehouse 选对；多个仓库时不要凭默认值操作。"]:
        add_bullet(doc, "□ " + s)
    doc.add_heading("下班前 1 分钟", level=2)
    for s in ["确认没有未同步的 PDA 离线记录。", "把未完成任务的单号和进度交接给下一班。", "右上角头像 → Sign out。", "公共电脑关闭浏览器窗口。"]:
        add_bullet(doc, "□ " + s)

    page_break(doc)
    section_title(doc, "2", "查单与筛选", "找不到记录时，90% 是筛选条件没清干净。")
    add_steps(doc, [
        ("先准备完整单号：例如 IB、OB、Container、Lot No.。", "复制时没有多余空格"),
        ("进入对应页面，在 Search 或对应编号框粘贴。", "筛选框里显示完整单号"),
        ("点击 Search。", "表格缩小到目标记录"),
        ("找不到时点击 Reset，再只输入一个最确定的编号。", "日期、状态、仓库等旧筛选被清空"),
        ("仍找不到，点 Refresh 后再查一次。", "页面重新加载"),
    ])
    add_callout(doc, "Global Search", "顶部 Global Search 适合用完整业务单号快速定位。模糊词可能出现多条结果，必须核对业务类型、仓库和状态后再打开。", "info")
    add_table(doc, ["现象", "先做什么", "不要做什么"], [
        ("列表空白", "Reset → Refresh → 重新搜索", "不要立刻新建同名记录"),
        ("重复记录", "核对编号、创建时间、状态", "不要随便取消其中一条"),
        ("按钮看不到", "确认账号权限，找主管", "不要借用别人账号"),
        ("页面卡住", "等待 10 秒，刷新并按编号核对", "不要连续点击 Save"),
    ], [2000, 3700, 3660])

    page_break(doc)
    section_title(doc, "3", "收货：创建 Inbound")
    doc.add_paragraph("菜单：InBound → Receiving")
    add_steps(doc, [
        ("点击 Create Inbound。", "右侧/弹窗出现新建表单"),
        ("填写 Container Number；选择 Customer、Warehouse。Warehouse 为必填。", "柜号、客户、仓库与纸面/通知一致"),
        ("按现场资料填写 Location、FC Code、Marking、Pallet Qty、Carton Qty、Weight LBS、CBM。", "数字单位没有混淆：PLT 托盘、CTN 箱、LBS 磅"),
        ("选择当前真实 Status；不确定时保留初始状态并问主管。", "状态没有超前"),
        ("点击 Save 一次。", "出现成功提示，列表里能搜到该 Inbound"),
    ])
    add_callout(doc, "最容易错的三项", "Container Number、Warehouse、数量单位。保存前把这三项和原始文件逐字核对。", "warn")
    doc.add_heading("收货状态怎么理解", level=2)
    add_table(doc, ["状态", "中文理解", "什么时候使用"], [
        ("Pending", "待处理", "尚未开始卸货"),
        ("Unloading", "卸货中", "现场正在卸货"),
        ("Received", "已收货", "收货数量已确认"),
        ("Put Away", "已上架", "货物已放到指定库位"),
        ("Completed", "已完成", "整票收货流程结束"),
        ("Hold", "暂停", "存在问题，等待处理"),
    ], [1800, 2000, 5560])
    add_callout(doc, "状态不能抢跑", "系统状态必须跟现场实物一致。货还在月台时，不要为了“先做完”而选 Put Away 或 Completed。", "danger")

    page_break(doc)
    section_title(doc, "4", "收货转库存：Receive to Inventory", break_before=False)
    doc.add_paragraph("这一步会创建库存 Lot，是收货与库存之间的关键动作。")
    add_steps(doc, [
        ("在 Receiving 页面搜索目标 Inbound。", "柜号、客户、仓库正确"),
        ("确认状态为 Received 或 Put Away。", "行内出现 Receive to Inventory"),
        ("再次核对实收托盘/箱数、库位和柜号。", "系统数据与现场点数一致"),
        ("点击 Receive to Inventory 一次。", "按钮变为 View Inventory 或显示已创建库存"),
        ("点击 View Inventory，核对 Lot No.、库位和数量。", "Inventory 页面能查到对应库存批次"),
    ])
    add_callout(doc, "看不到按钮", "只有 Received / Put Away 状态才会出现 Receive to Inventory。先确认状态；如果无权修改或数据不对，找主管，不要新建另一票 Inbound。", "info")
    add_callout(doc, "防重复", "如果已经显示 View Inventory，说明库存已创建。不要再次用导入或调整的方式“补一份”。", "danger")

    page_break(doc)
    section_title(doc, "5", "库存查询与四种操作")
    doc.add_paragraph("菜单：Warehouse → Inventory。优先用 Lot No.、Container 或 FC Code 搜索。")
    add_table(doc, ["操作", "意思", "使用前必须确认"], [
        ("Move", "移动库位", "实物已经/马上会移动，New Location 正确"),
        ("Adjust", "盘点调整", "真实差异已复核；Reason 必填；通常需主管授权"),
        ("Hold", "冻结库存", "这部分暂时不能被正常使用"),
        ("Release", "解除冻结", "问题已解决，可以恢复使用"),
    ], [1400, 2500, 5460])
    doc.add_heading("通用操作方法", level=2)
    add_steps(doc, [
        ("搜索并打开目标 Lot，核对 Container、Warehouse、Location。", "目标批次唯一且正确"),
        ("点击 More / …，选择 Move、Adjust、Hold 或 Release。", "弹出对应表单"),
        ("填写新库位/数量/原因/备注。Adjust 的 OTHER 必须写 Remark。", "调整后数量不会小于 0"),
        ("请第二个人或主管复核关键数据。", "实物、单据、系统三者一致"),
        ("提交一次。", "出现 “Inventory operation completed”"),
    ])
    add_callout(doc, "操作被拒绝", "看到 “Operation rejected. Check quantities, location, and permissions.”：检查数量是否超出、库位是否有效、账号是否有权限。不要为了通过而随意改小数量。", "warn")
    add_callout(doc, "库存数量关系", "Available（可用）可能小于 On Hand（在库），因为部分库存已分配或冻结。不要把“不可用”误认为“丢失”。", "info")

    page_break(doc)
    section_title(doc, "6", "柜况与优先级：Container Tracking", break_before=False)
    doc.add_paragraph("菜单：InBound → Container Tracking。点击 Track / 查柜，可查看同一柜关联的收货、库存、FBA 和出库任务。")
    add_steps(doc, [
        ("输入完整 Container Number。", "柜号格式与原始文件一致"),
        ("点击 Track / 查柜。", "打开柜况详情"),
        ("查看 Earliest Outbound、Days、Priority、Readiness。", "知道最早出库日期和是否具备出库条件"),
        ("若显示 NOT_READY，展开关联数据检查库存或分配。", "找到缺库存或无有效分配的原因"),
    ])
    add_table(doc, ["Priority", "含义", "建议"], [
        ("CRITICAL", "已逾期或今天到期", "立即报告并优先处理"),
        ("HIGH", "1–2 天内到期", "列入当前班次重点"),
        ("MEDIUM", "3–5 天内到期", "按计划处理"),
        ("NORMAL", "更晚或未排期", "正常队列"),
    ], [1700, 3300, 4360])
    add_callout(doc, "Readiness 不是实物确认", "READY 只表示系统条件满足；装车、封柜、交接仍要按现场流程确认。NOT_READY 时不要强行推进状态。", "warn")

    page_break(doc)
    section_title(doc, "7", "拣货单与打印")
    doc.add_paragraph("菜单：OutBound → Picking List。已完成记录可在 Picking History 查看。")
    add_steps(doc, [
        ("在 Picking List 找到主管分配的拣货单。", "Picking No.、仓库、状态正确"),
        ("点击 Print，打印或保存拣货单。", "纸面单号和系统一致"),
        ("按单据到指定库位拣货；需要扫码时进入 PDA。", "库位、Lot No.、数量三项一致"),
        ("全部拣完并复核后，按岗位要求点击 Complete。", "出现 “Picking completed”"),
        ("到 Picking History 按单号确认。", "记录已进入历史列表"),
    ])
    add_callout(doc, "重要业务规则", "Picking Complete 只完成拣货任务，不会直接扣减库存。最终库存扣减由 Outbound Complete 负责。不要因为库存暂时没变而重复完成。", "info")
    add_callout(doc, "不要扫错码", "系统识别 Picking No.、Location Code 和 Inventory Lot No.。托盘/LPN/SKU 条码不是本流程的替代码。", "danger")

    page_break(doc)
    section_title(doc, "8", "PDA 扫码拣货", "严格按“库位 → 批次 → 数量”循环，不要跳步。")
    doc.add_heading("进入作业", level=2)
    add_steps(doc, [
        ("打开 PDA 页面，选择作业仓库。", "右上角显示在线/离线状态"),
        ("扫描或输入拣货单号。", "显示正确的 Picking No."),
        ("点击 进入作业。", "看到进度和待拣明细"),
    ])
    doc.add_heading("每一笔都这样做", level=2)
    add_steps(doc, [
        ("第1步：扫描库位 Location。", "页面接受库位并进入第2步"),
        ("第2步：扫描库存批次 Lot No.。", "页面显示当前库位与批次"),
        ("第3步：输入本次实际拣货数量，点击 确认数量 / 提交扫描。", "最近记录显示 ACCEPTED，已拣数量增加"),
        ("换库位或批次时，从第1步重新扫描。", "当前上下文清空后重新开始"),
    ])
    add_callout(doc, "数量规则", "数量必须大于 0，且不能超过剩余待拣数量。系统拒绝时先核对实物和剩余数量，不要拆成多次乱试。", "warn")
    doc.add_heading("结束作业", level=2)
    add_table(doc, ["按钮", "什么时候点"], [
        ("完成并结束会话", "全部拣完、最近记录无未处理错误、复核完成"),
        ("结束本次作业（部分完成）", "换班、暂停或现场无法继续；必须交接剩余数量"),
        ("重新扫描库位", "扫错库位/批次，但尚未确认数量"),
    ], [3300, 6060])

    page_break(doc)
    section_title(doc, "9", "PDA 离线时怎么办", break_before=False)
    add_callout(doc, "先认清提示", "在线时按钮通常是“提交扫描”；离线时会显示“保存到本机”。保存到本机 ≠ 服务器已收到。", "danger")
    add_steps(doc, [
        ("继续前先确认页面明确显示离线保存可用。", "有离线提示且本地保存正常"),
        ("按正常顺序扫描；每笔确认后查看最近记录。", "记录保存在本机队列"),
        ("不要关闭浏览器、清缓存、换设备或退出账号。", "离线队列仍保留"),
        ("网络恢复后保持页面打开，等待自动同步。", "看到同步成功/队列清空"),
        ("按 Picking No. 核对服务器进度后再结束。", "已拣数量与现场一致"),
    ])
    add_callout(doc, "立即停止的情况", "页面提示无法保存到本机、队列长期不同步、同一笔反复失败、设备要关机或换人。记下 Picking No. 和未同步笔数，联系主管。", "danger")
    add_table(doc, ["现象", "处理"], [
        ("离线记录显示待同步", "保持页面打开，恢复网络后等待"),
        ("同步失败", "截图错误；不要删除队列；找主管"),
        ("不确定是否上传", "服务器端 Refresh，按 Picking No. 核对进度"),
        ("本地无法保存", "停止扫码，改用主管批准的备用流程"),
    ], [3200, 6160])

    page_break(doc)
    section_title(doc, "10", "出库 Outbound：一线员工怎么配合", break_before=False)
    doc.add_paragraph("菜单：OutBound → Dispatch。Create OB、Allocate、Confirm、Dispatch、Complete、Cancel 属于关键业务动作，只由被授权岗位执行。")
    add_table(doc, ["状态/动作", "中文理解", "现场要求"], [
        ("New", "新建", "资料可能尚未齐全"),
        ("In Progress", "处理中", "正在分配或准备"),
        ("Confirmed", "已确认", "出库资料与分配已确认"),
        ("Dispatched", "已发运", "车辆/货物已按公司流程交接"),
        ("Completed", "已完成", "最终完成；会承担库存最终扣减"),
        ("Exception", "异常", "先处理异常，不要继续硬推"),
        ("Canceled", "已取消", "业务作废"),
    ], [2000, 2600, 4760])
    add_callout(doc, "点状态前四核对", "OB#、Warehouse、分配来源/数量、现场实物。任何一项不一致，就不要 Confirm / Dispatch / Complete。", "danger")
    add_callout(doc, "常见错误", "看到 “Invalid status or missing required data” 说明状态顺序不对或资料不完整；看到分配超量提示，说明来源可用量不足。不要修改实物数量来迎合系统。", "warn")
    doc.add_heading("FBA Workbench", level=2)
    doc.add_paragraph("FBA 页面用于批量出库协作。Create FBA、Import Excel、Batch Picking / BOL / Outbound 只按主管安排使用。批量操作前先确认选中的每一行，避免把无关任务一起处理。")

    page_break(doc)
    section_title(doc, "11", "BOL、Loads 与发运资料")
    doc.add_heading("BOL", level=2)
    add_steps(doc, [
        ("菜单 OutBound → BOL，按单号找到目标记录。", "单号、承运商、目的地一致"),
        ("下载 PDF 或 Excel。", "文件成功打开，内容无空白/乱码"),
        ("打印前再次核对数量和地址。", "纸面与系统一致"),
    ])
    doc.add_heading("Loads", level=2)
    add_steps(doc, [
        ("菜单 OutBound → Loads，搜索 Load No.。", "目标 Load 唯一"),
        ("打开详情，查看 Outbound Orders、Work Orders、Documents 和 Active Exceptions。", "没有未处理关键异常"),
        ("按主管要求推进 PLANNED → READY → DISPATCHED → COMPLETED。", "状态与现场实际同步"),
    ])
    add_callout(doc, "发运前最后检查", "车辆/承运商、Load No.、OB#、托盘/箱数、BOL、异常状态。发现 Active Exception 时先问主管是否允许继续。", "warn")

    page_break(doc)
    section_title(doc, "12", "Work Orders：接任务与报进度")
    doc.add_paragraph("菜单：Operations → Work Orders。常见类型：PICK、STAGE、LOAD、CHECK、GENERAL。")
    add_table(doc, ["状态", "你应该怎么做"], [
        ("OPEN", "待分配；不要自行抢不属于你的任务"),
        ("ASSIGNED", "已分配；核对负责人、仓库、关联业务单"),
        ("IN_PROGRESS", "已开始；现场确实开工后再改"),
        ("COMPLETED", "任务全部完成且结果已记录"),
        ("CANCELED", "任务取消；停止继续操作并确认交接"),
    ], [2100, 7260])
    add_steps(doc, [
        ("搜索 Work Order No.，打开详情。", "类型、Priority、Warehouse、关联单号正确"),
        ("确认 Assignment 是你或你的班组。", "职责明确"),
        ("开始现场工作后改为 IN_PROGRESS。", "状态更新并留在历史记录"),
        ("在 Notes / Documents 补充必要信息。", "同事能看懂发生了什么"),
        ("完成并复核后改为 COMPLETED。", "没有遗漏步骤或未上传证据"),
    ])

    page_break(doc)
    section_title(doc, "13", "异常处理：Trouble Shoot")
    doc.add_paragraph("菜单：Operations → Trouble Shoot。异常不是“丢脸”，及时、准确记录比私下修补更重要。")
    add_steps(doc, [
        ("点击 Create Exception。", "打开异常表单"),
        ("选择类型和严重程度，写清发生了什么。", "描述包含时间、地点、数量、影响"),
        ("至少关联一个真实业务记录。", "Business Reference 能打开对应单据"),
        ("提交后记下 Exception No.。", "提示 Created + 异常号"),
        ("按负责人推进 OPEN → INVESTIGATING → RESOLVED。", "解决说明 Resolution 完整"),
    ])
    add_callout(doc, "描述模板", "“在【时间】、【仓库/库位】，处理【业务单号】时发现【问题】；涉及【数量】；目前已【暂停/隔离/保留现场】；需要【谁】协助。”", "info")
    add_callout(doc, "不要直接取消", "Cancel 表示异常记录作废，不等于问题解决。只有确认误报或按主管要求时才取消。需要现场复核时可创建 CHECK / HIGH Work Order。", "warn")

    page_break(doc)
    section_title(doc, "14", "Documents & POD：上传凭证")
    doc.add_paragraph("菜单：Operations → Documents & POD。常见类型：BOL、POD、DELIVERY_RECEIPT、WAREHOUSE、EXCEPTION_ATTACHMENT、GENERAL。")
    add_steps(doc, [
        ("点击 Upload，选择正确的 Document Type。", "类型与文件内容一致"),
        ("选择文件：PDF、Excel、CSV、JPG/JPEG 或 PNG。", "文件能正常打开，方向清楚"),
        ("至少关联一个业务记录：Load、Outbound、Work Order、Exception 或 Container。", "关联编号准确"),
        ("填写必要说明后上传。", "列表出现新文件，可 Download 打开"),
    ])
    add_callout(doc, "拍照标准", "四角完整、文字清楚、没有手指遮挡、方向正确；一张图只表达一个主要凭证。客户签字或敏感信息按公司规定处理。", "success")
    add_callout(doc, "Archive 不是普通关闭", "Archive 会把文档归档隐藏，属于管理动作。传错文件时先保留证据并联系主管，不要自行反复上传、归档。", "warn")

    page_break(doc)
    section_title(doc, "15", "常见问题：照着排查")
    add_table(doc, ["问题", "第一步", "第二步", "仍不行"], [
        ("登录失败", "检查大小写", "检查网络，只重试一次", "截图找管理员"),
        ("找不到单", "Reset", "只用完整单号搜索 + Refresh", "确认仓库/权限"),
        ("保存后没反应", "等待 10 秒", "Refresh 后按单号核对", "不要重复点，找主管"),
        ("按钮是灰色", "看当前状态", "核对前置步骤和权限", "找主管"),
        ("库存操作被拒", "核对数量/库位", "核对权限和可用量", "截图错误"),
        ("PDA 扫码被拒", "核对库位→Lot 顺序", "核对剩余数量", "停止并报告"),
        ("PDA 离线不同步", "保持页面打开", "恢复网络等待队列清空", "保留队列，找主管"),
        ("导出文件找不到", "看浏览器下载图标", "打开电脑“下载”文件夹", "重新导出一次"),
        ("页面数据像旧的", "Refresh", "退出再登录", "报告时间和页面"),
    ], [2100, 2200, 2700, 2360])
    doc.add_heading("向主管报告时，一次说全", level=2)
    for s in ["你的姓名/班组", "发生时间", "页面名称", "业务单号", "你刚做的最后一步", "完整错误提示或截图", "现场实物当前状态", "是否存在未同步离线记录"]:
        add_bullet(doc, "□ " + s)

    page_break(doc)
    section_title(doc, "附录 A", "管理员导入：只给授权人员")
    add_callout(doc, "普通员工可跳过", "Import Excel / Data Upload 会批量改变数据。没有培训和权限时不要使用。", "danger")
    add_steps(doc, [
        ("优先点击 Download Template 获取当前模板。", "使用系统最新字段"),
        ("准备 .xlsx 或 UTF-8 .csv；不要使用旧 .xls。", "文件可正常打开，表头完整"),
        ("上传并完成字段 Mapping。", "必填字段全部映射"),
        ("先看 Validation / Preview。", "错误行已修正；Preview 不会正式写入"),
        ("确认 Duplicate 处理规则；默认优先 SKIP。", "知道跳过或更新的后果"),
        ("最后 Confirm，再查看 Result 和 Import History。", "成功/跳过/失败数量可解释"),
    ])
    add_callout(doc, "导入纪律", "不要为了消除报错而随意改业务值；不要把 UPDATE 当默认选项；导入结果必须留存文件名、时间、操作者和错误行。", "warn")

    page_break(doc)
    section_title(doc, "附录 B", "班次快速检查卡")
    add_table(doc, ["阶段", "检查项目", "完成"], [
        ("上班", "本人账号、网络、正确仓库、刷新任务", "□"),
        ("收货", "柜号、仓库、PLT/CTN、真实状态", "□"),
        ("入库", "Received/Put Away、库位、Receive to Inventory 防重复", "□"),
        ("库存", "Lot、Location、Available、Hold/Allocation", "□"),
        ("拣货", "Picking No.、库位、Lot、实际数量", "□"),
        ("PDA", "最近记录 ACCEPTED、离线队列已同步", "□"),
        ("出库", "OB、Load、BOL、实物、异常状态", "□"),
        ("交接", "未完成单号、剩余数量、异常号、退出账号", "□"),
    ], [1500, 6760, 1100])
    add_callout(doc, "最后一句", "系统是记录现场事实的工具。先确认实物，再操作系统；不确定就停手、留证、找主管。", "success")

    # Apply consistent table geometry and paragraph controls.
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        tbl_layout = tbl_pr.find(qn("w:tblLayout"))
        if tbl_layout is None:
            tbl_layout = OxmlElement("w:tblLayout")
            tbl_pr.append(tbl_layout)
        tbl_layout.set(qn("w:type"), "fixed")
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        if tbl_ind is None:
            tbl_ind = OxmlElement("w:tblInd")
            tbl_pr.append(tbl_ind)
        tbl_ind.set(qn("w:w"), "120")
        tbl_ind.set(qn("w:type"), "dxa")
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    set_paragraph_keep(p)

    props = doc.core_properties
    props.title = "DLX Yuki WMS 员工操作手册（零基础版）"
    props.subject = "仓库员工日常操作、扫码拣货、异常处理与单据归档"
    props.author = "DLX Yuki WMS"
    props.keywords = "WMS, 员工手册, 收货, 库存, 拣货, 出库, PDA"
    props.comments = "根据当前系统菜单与操作逻辑编制。"

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
