#!/usr/bin/env python3
"""Offline converter for 美西仓 4.0 → Yuki WMS import CSVs.

Does not write to the database. Use START_DLX_WMS.bat for the live system.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

OL_HEADERS = ["柜号", "仓点", "库位", "板数", "件数", "磅数(lb)", "体积", "客户", "拆柜时间", "实际到仓时间", "FBA", "PO", "备注"]
DS_HEADERS = ["客户", "柜号", "ID / ST", "FBA", "PO", "件数", "LBS", "体积", "派送仓点", "备注", "关联OL"]
OUTBOUND_HEADERS = [
    "计划单号", "仓点", "ISA/预约编号", "预计送仓时间", "出库时间", "真实板数", "拣货单", "BOL",
    "预计体积", "预计板数", "重量(lb)", "车队", "Redirect Code", "预计件数", "POD", "PO", "FBA",
    "真实件数", "真实重量", "真实体积", "计划出库",
]
TRACKING_HEADERS = [
    "container_number", "mbl_number", "pod_eta", "actual_delivery_at",
    "wa_received_at", "wa_complete_at", "customer_reference",
    "container_attributes", "container_remark", "source_row_number",
]

CONTAINER_RE = re.compile(r"^[A-Z]{4}\d{7}$")


def norm_header(value: Any) -> str:
    return " ".join(str(value or "").strip().replace("／", "/").split())


def iso(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def num(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError:
        return 0.0


def container(value: Any) -> str:
    return str(value or "").strip().upper().replace(" ", "")


def pick(row: dict[str, Any], *names: str) -> Any:
    lowered = {norm_header(k).casefold(): v for k, v in row.items()}
    for name in names:
        key = norm_header(name).casefold()
        if key in lowered and lowered[key] not in (None, ""):
            return lowered[key]
    return None


def open_book(path: Path):
    return load_workbook(path, read_only=True, data_only=True, keep_links=False)


def headers_of(ws) -> list[str]:
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not first:
        return []
    used: dict[str, int] = {}
    out: list[str] = []
    for i, value in enumerate(first, start=1):
        base = str(value).strip() if value not in (None, "") else f"Column {i}"
        used[base] = used.get(base, 0) + 1
        out.append(base if used[base] == 1 else f"{base}_{used[base]}")
    return out


def iter_sheet(path: Path, sheet_name: str):
    wb = open_book(path)
    try:
        if sheet_name not in wb.sheetnames:
            return
        ws = wb[sheet_name]
        headers = headers_of(ws)
        for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
            if any(v not in (None, "") for v in row.values()):
                yield row_number, row
    finally:
        wb.close()


def inspect(path: Path) -> dict[str, Any]:
    wb = open_book(path)
    try:
        sheets = list(wb.sheetnames)
        estimates = {name: max(0, int(wb[name].max_row or 0) - 1) for name in sheets}
    finally:
        wb.close()

    tigui_name = next((n for n in sheets if n in ("提柜", "柜", "Container")), None)
    containers_tigui: list[str] = []
    dupes: Counter[str] = Counter()
    if tigui_name:
        seen: Counter[str] = Counter()
        for _, row in iter_sheet(path, tigui_name):
            no = container(pick(row, "柜号", "container_number", "Container"))
            if no:
                seen[no] += 1
                containers_tigui.append(no)
        dupes = Counter({k: v for k, v in seen.items() if v > 1})

    ol_by_c: dict[str, dict[str, float]] = defaultdict(lambda: {"rows": 0, "pallet": 0, "carton": 0, "lbs": 0, "cbm": 0})
    loc_values: Counter[str] = Counter()
    customers: Counter[str] = Counter()
    if "OL" in sheets:
        for _, row in iter_sheet(path, "OL"):
            no = container(pick(row, "柜号"))
            if not no:
                continue
            rec = ol_by_c[no]
            rec["rows"] += 1
            rec["pallet"] += num(pick(row, "板数"))
            rec["carton"] += num(pick(row, "件数"))
            rec["lbs"] += num(pick(row, "磅数(lb)", "磅数", "磅数 (lb)"))
            rec["cbm"] += num(pick(row, "体积", "方数"))
            loc = str(pick(row, "仓点") or "").strip()
            if loc:
                loc_values[loc] += 1
            cust = str(pick(row, "客户") or "").strip()
            if cust:
                customers[cust] += 1

    outbound_by_plan = 0
    if "出库" in sheets:
        outbound_by_plan = sum(1 for _ in iter_sheet(path, "出库"))

    ds_by_c: Counter[str] = Counter()
    if "DS" in sheets:
        for _, row in iter_sheet(path, "DS"):
            no = container(pick(row, "柜号"))
            if no:
                ds_by_c[no] += 1

    both = sorted(set(ol_by_c) & set(containers_tigui or ol_by_c))
    return {
        "file": str(path),
        "size_mb": round(path.stat().st_size / 1024 / 1024, 1),
        "sheets": sheets,
        "row_estimates": estimates,
        "tigui_sheet": tigui_name,
        "tigui_containers": len(set(containers_tigui)),
        "tigui_rows": len(containers_tigui),
        "duplicate_containers": dict(dupes),
        "ol_containers": len(ol_by_c),
        "ol_rows": int(sum(v["rows"] for v in ol_by_c.values())),
        "ds_rows_with_container": int(sum(ds_by_c.values())),
        "outbound_rows": outbound_by_plan,
        "location_values": len(loc_values),
        "customers": len(customers),
        "location_top": loc_values.most_common(20),
        "overlap_tigui_ol": len(both),
    }


def choose_pilot(path: Path, count: int, named: list[str] | None) -> list[str]:
    if named:
        return [container(x) for x in named if container(x)]

    tigui_done: set[str] = set()
    wb = open_book(path)
    try:
        wb_sheets = list(wb.sheetnames)
    finally:
        wb.close()
    tigui_name = next((n for n in wb_sheets if n in ("提柜", "柜")), None)
    if tigui_name:
        for _, row in iter_sheet(path, tigui_name):
            no = container(pick(row, "柜号"))
            devan = pick(row, "完成拆柜日期", "拆柜完成", "wa_complete_at", "拆柜时间")
            if no and devan not in (None, ""):
                tigui_done.add(no)

    ol_rows: dict[str, int] = Counter()
    ol_locs: dict[str, set[str]] = defaultdict(set)
    ol_cartons: dict[str, float] = defaultdict(float)
    if "OL" in wb_sheets:
        for _, row in iter_sheet(path, "OL"):
            no = container(pick(row, "柜号"))
            if not no:
                continue
            ol_rows[no] += 1
            ol_locs[no].add(str(pick(row, "仓点") or ""))
            ol_cartons[no] += num(pick(row, "件数"))

    ds_set: set[str] = set()
    if "DS" in wb_sheets:
        for _, row in iter_sheet(path, "DS"):
            no = container(pick(row, "柜号"))
            if no:
                ds_set.add(no)

    pool = [c for c in ol_rows if c in tigui_done] or list(ol_rows)
    simple = [c for c in pool if len(ol_locs[c]) <= 1]
    multi = [c for c in pool if len(ol_locs[c]) > 1]
    shipped = [c for c in pool if c in ds_set]
    large = sorted(pool, key=lambda c: ol_cartons[c], reverse=True)
    picked: list[str] = []

    def add(seq: list[str], n: int) -> None:
        for item in seq:
            if item not in picked:
                picked.append(item)
            if len(picked) >= n:
                return

    add(simple, 2)
    add(multi, len(picked) + 2)
    add(shipped, len(picked) + 2)
    add(large, len(picked) + 1)
    add(pool, count)
    return picked[:count]


def write_csv(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({h: row.get(h, "") for h in headers})


def map_tracking(row_number: int, row: dict[str, Any]) -> dict[str, Any]:
    return {
        "container_number": container(pick(row, "柜号")),
        "mbl_number": pick(row, "MBL", "MBL号", "提单号", "主提单") or "",
        "pod_eta": iso(pick(row, "ETA", "预计到港", "预计到仓")),
        "actual_delivery_at": iso(pick(row, "实际到仓日期", "实际到仓时间", "到仓日期")),
        "wa_received_at": iso(pick(row, "实际到仓日期", "实际到仓时间")),
        "wa_complete_at": iso(pick(row, "完成拆柜日期", "拆柜完成", "拆柜时间")),
        "customer_reference": pick(row, "客户") or "",
        "container_attributes": pick(row, "柜型", "箱型") or "",
        "container_remark": pick(row, "备注") or "",
        "source_row_number": row_number,
    }


def project_row(row: dict[str, Any], wanted: list[str]) -> dict[str, Any]:
    lowered = {norm_header(k).casefold(): ("" if v is None else v) for k, v in row.items()}
    out: dict[str, Any] = {}
    for header in wanted:
        key = norm_header(header).casefold()
        value = lowered.get(key, "")
        if header == "柜号":
            value = container(value)
        elif header in ("拆柜时间", "实际到仓时间", "预计送仓时间", "出库时间", "ETA"):
            value = iso(value)
        out[header] = value
    return out


def export(path: Path, out_dir: Path, containers: list[str]) -> dict[str, Any]:
    wanted = set(containers)
    wb = open_book(path)
    try:
        sheets = list(wb.sheetnames)
    finally:
        wb.close()

    tigui_name = next((n for n in sheets if n in ("提柜", "柜")), None)
    tracking_rows: list[dict[str, Any]] = []
    if tigui_name:
        seen: set[str] = set()
        for row_number, row in iter_sheet(path, tigui_name):
            mapped = map_tracking(row_number, row)
            no = mapped["container_number"]
            if no and no in wanted and no not in seen:
                tracking_rows.append(mapped)
                seen.add(no)

    ol_rows: list[dict[str, Any]] = []
    ol_sum: dict[str, dict[str, float]] = defaultdict(lambda: {"rows": 0, "pallet": 0, "carton": 0, "lbs": 0, "cbm": 0})
    if "OL" in sheets:
        for _, row in iter_sheet(path, "OL"):
            no = container(pick(row, "柜号"))
            if no not in wanted:
                continue
            projected = project_row(row, OL_HEADERS)
            ol_rows.append(projected)
            rec = ol_sum[no]
            rec["rows"] += 1
            rec["pallet"] += num(projected.get("板数"))
            rec["carton"] += num(projected.get("件数"))
            rec["lbs"] += num(projected.get("磅数(lb)"))
            rec["cbm"] += num(projected.get("体积"))

    ds_rows: list[dict[str, Any]] = []
    if "DS" in sheets:
        for _, row in iter_sheet(path, "DS"):
            no = container(pick(row, "柜号"))
            if no in wanted:
                ds_rows.append(project_row(row, DS_HEADERS))

    outbound_rows: list[dict[str, Any]] = []
    if "出库" in sheets:
        for _, row in iter_sheet(path, "出库"):
            no = container(pick(row, "柜号"))
            if no in wanted:
                outbound_rows.append(project_row(row, OUTBOUND_HEADERS))

    write_csv(out_dir / "01_container_tracking.csv", TRACKING_HEADERS, tracking_rows)
    write_csv(out_dir / "02_ol_inbound.csv", OL_HEADERS, ol_rows)
    write_csv(out_dir / "03a_outbound_header.csv", OUTBOUND_HEADERS, outbound_rows)
    write_csv(out_dir / "03b_ds_lines.csv", DS_HEADERS, ds_rows)

    reconcile = []
    for no in containers:
        rec = ol_sum.get(no, {"rows": 0, "pallet": 0, "carton": 0, "lbs": 0, "cbm": 0})
        bad_format = bool(no) and CONTAINER_RE.match(no) is None
        reconcile.append({
            "container": no,
            "tracking": any(r["container_number"] == no for r in tracking_rows),
            "ol_rows": rec["rows"],
            "pallet": rec["pallet"],
            "carton": rec["carton"],
            "lbs": round(rec["lbs"], 3),
            "cbm": round(rec["cbm"], 4),
            "ds_rows": sum(1 for r in ds_rows if r.get("柜号") == no),
            "outbound_rows": sum(1 for r in outbound_rows if container(r.get("柜号")) == no),
            "iso_container": not bad_format,
        })
    (out_dir / "reconcile.json").write_text(json.dumps({
        "source": str(path),
        "containers": containers,
        "files": {
            "tracking": len(tracking_rows),
            "ol": len(ol_rows),
            "outbound": len(outbound_rows),
            "ds": len(ds_rows),
        },
        "rows": reconcile,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"out_dir": str(out_dir), "containers": containers, "counts": {
        "tracking": len(tracking_rows), "ol": len(ol_rows),
        "outbound": len(outbound_rows), "ds": len(ds_rows),
    }}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Convert 美西仓 4.0 workbook into Yuki import CSVs (no database writes)")
    sub = parser.add_subparsers(dest="cmd", required=True)
    p_ins = sub.add_parser("inspect", help="Read-only sheet / container summary")
    p_ins.add_argument("workbook", type=Path)
    p_exp = sub.add_parser("export", help="Write pilot or named-container CSVs")
    p_exp.add_argument("workbook", type=Path)
    p_exp.add_argument("--out", type=Path, default=Path("import_out"))
    p_exp.add_argument("--pilot", type=int, default=8)
    p_exp.add_argument("--containers", type=str, default="", help="Comma-separated container numbers")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    workbook: Path = args.workbook.expanduser().resolve()
    if not workbook.is_file():
        print(f"Workbook not found: {workbook}", file=sys.stderr)
        return 2
    if args.cmd == "inspect":
        print(json.dumps(inspect(workbook), ensure_ascii=False, indent=2))
        return 0
    named = [part.strip() for part in args.containers.split(",") if part.strip()]
    containers = choose_pilot(workbook, args.pilot, named or None)
    if not containers:
        print("No containers selected", file=sys.stderr)
        return 3
    print(json.dumps(export(workbook, args.out.resolve(), containers), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
